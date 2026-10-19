"""Build a journal article about the final Bora method and its fixed evaluation.

Uses completed evidence only. Does not train, infer, submit jobs or submit a paper.
Authoring requires python-docx and Pillow in the bundled document runtime.
"""
import argparse
import json
from pathlib import Path
import re
from statistics import median
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

TITLE = 'Bora for calibrated refinement and identification of tissue regions in histopathology'
ABSTRACT = [
('Background and Objective', 'Coarse tile maps obscure tissue boundaries. We developed Bora to refine supplied partitions, identify complete tissue groups and export compact native-coordinate polygons.'),
('Methods', 'Physically calibrated Virchow2 evidence guides H&E appearance fields and categorical competition. Neighbor context combines group feature means with physical shared-boundary-weighted neighborhood contrasts in one fastPLS classifier. Shared-boundary fitting and polygon/raster checks remove fragments and holes below 20 native pixels of area. We assessed 39 matched regions and 116 multiclass images.'),
('Results', 'On 39 matched regions, mean macro intersection-over-union was 0.663 after refinement, versus 0.586 for coarse tiles and 0.643 for a bilateral reference. On 116 development images, Neighbor context classified 350 of 398 groups correctly, versus 340 for a pooled-mean control, and achieved mean final-polygon intersection-over-union of 0.587 versus 0.561. The originally designated test and validation subsets had means of 0.597 and 0.566, respectively; both were previously inspected. Tumor–stroma interface F1 was 0.571 at 8 µm tolerance, with nine missed interfaces. Recorded cached deployment took 2.071 seconds per image. An exploratory 3.67-billion-pixel whole-slide run completed exact native export but missed normal glands. In one source-to-target ROI pilot, macro IoU increased from 0.603 for predicted tiles to 0.771 after Bora.'),
('Conclusions', 'Bora supports efficient conditional refinement and tissue identification. Expert-derived inputs and previously inspected evaluation images require interpreting these findings as exploratory. Independent validation with predicted inputs remains necessary.')]

SECTIONS = [
('Introduction', [
'Tissue compartments such as tumor, stroma and necrosis are central to quantitative pathology. Foundation models provide informative tile representations [1], but one assignment per tile cannot directly describe irregular tissue interfaces. Tissue identity and tissue geometry also operate at different scales: several mixed tiles may contribute to one region, while disconnected pieces may share a tissue identity. A practical analysis tool should refine a supplied partition and assign interpretable identities to complete tissue groups.',
'Bora addresses this conditional problem. It accepts an aligned H&E image, anonymous tile assignments, calibrated Virchow2 spatial features and a fitted tissue classifier, forms an image-guided partition and identifies each refined group using its appearance and neighborhood. Its output includes a categorical raster, compact GeoJSON with predicted tissue identities and an auditable link to physical image coordinates. The contribution is an integrated computational method for calibrated whole-group refinement and identification, rather than a new partial least-squares estimator.',
'Graph-based image segmentation, edge-aware fitting and local appearance models provide relevant precedents [2–4]. Feature upsampling methods address restoration of spatial detail in representations [5], while dense segmentation models learn direct image-to-label mappings. Bora instead retains an explicit coarse-map input and separates geometry from semantic identification. This distinction permits the same refined geometry to be inspected with or without named tissue labels, but also makes performance conditional on the upstream partition.'
' Recent Virchow2-to-student distillation and multi-resolution promptable segmentation address related computational costs and spatial detail through different input and training assumptions [11,12].',
]),
('Methods', []),
('Inputs and physical calibration', [
'Inputs comprise an H&E image, anonymous tile-cluster assignments as a grid or aligned categorical coarse mask, image-derived tissue support, Virchow2 spatial tokens, a fitted tissue classifier and native/working pixel spacing along both axes. Working images use approximately 0.5 µm per pixel; effective spacing after integer resizing is retained explicitly. A 224-pixel encoder window therefore covers approximately 112 µm, while a 14-pixel token interval corresponds to approximately 7 µm. These values describe sampling and context, not optical resolving power. Full 1280-channel spatial tokens are retained, with the encoder revision recorded in the reproducibility manifest [1].',
'The evaluation uses tissue-region annotations rather than nuclei or gland-instance endpoints. Fields with uncertain calibration, insufficient usable coarse observations or only one mapped tissue group are ineligible for multiclass evaluation. At least 16 usable coarse observations are required. Images are recorded as ROIs or WSI-derived windows; native coordinates do not imply that complete slides were processed.',
'For the controlled task, majority expert labels define approximately 45-µm coarse tile assignments. A deterministic permutation replaces their class numbers with anonymous group identifiers without changing assignments. The identity key supplies training and evaluation targets separately; source names and the per-image reference vocabulary do not enter the identifier. This preserves the expert-derived partition information and must not be interpreted as annotation-free segmentation.'
]),
('Virchow2 evidence conversion', [
'The fixed RGB conversion combines scale-persistent class evidence with predictable latent structure. Balanced coarse observations support a 64-coordinate minimum-redundancy selector. Scale-specific evidence is aligned at approximately 0.5, 1 and 2 µm per pixel, with 32 pair-balanced coordinates and a regularized linear classifier. A second term estimates predictable structure from four spatial quadrants through ridge fitting, using a four-token exclusion buffer and context scales of one, two and four. Their class-evidence fields are averaged before palette conversion.',
'The resulting three-channel image is an evidence representation, not reconstructed H&E, measured protein expression or a microscopy channel. Conversion uses the supplied coarse groups. Although the fine reference mask is not read during refinement, the RGB evidence remains expert-conditioned in this benchmark because those groups originate from expert annotations. Unannotated tissue is included in feature extraction and image processing; only its contribution to reference quality scores is excluded.'
]),
('H and E guided appearance fields', [
'Bora represents evidence on an approximately 1-µm analysis grid and fits one three-channel appearance field per group on an approximately 8-µm grid. The regularizer connects adjacent field sites. Edge conductance decreases with squared H&E optical-density difference, with contrast scale 0.1 and floor 0.02; physical spacing scales a 16-µm regularization length. Coarse observations provide soft class-specific data weights, and a prior of strength 0.015 anchors each field to its global class mean.',
'For group k, Dₖ contains its spatial observation masses, Bₖ contains weighted evidence sums, Lw is the H&E-guided physical graph Laplacian and μₖ is the global evidence mean. The field Fₖ is obtained from the linear system below. The positive prior makes the fixed-observation system nonsingular. The computational problem is a sparse quadratic fit for each class, not a learned dense decoder.',
'Global statistics use balanced capped coarse samples. The pooled covariance combines 80% empirical covariance, 20% diagonal covariance and a ridge of 0.0025 in standardized evidence units. Interpolated local fields define negative Mahalanobis residuals over supported tissue. Softmax competition yields relative scores, and the maximum selects the group. No fixed boundary margin or immutable label interior restricts changes. These scores are not calibrated posterior tissue probabilities.',
'The categorical result is projected directly to the requested native grid. H&E-derived exterior background is removed afterward. Errors in tissue support therefore remain possible, including loss of pale annotated tissue. Unannotated tissue is not assigned a negative reference label.'
]),
('Neighbor context tissue identification', [
'The prediction unit is one anonymous tissue group, including all disconnected pieces carrying its identifier. Native-pixel overlap weights pool the full 1280-channel spatial tokens into a group mean μg. This avoids treating disconnected fragments as independent training examples. Pooling respects the token grid and physical projection; padded token borders are not stretched over the image.',
'For each adjacent group h, ℓgh is the shared interface length in micrometres. The neighborhood mean νg is the length-weighted average of adjacent group means. The descriptor concatenates the group mean and its contrast to that neighborhood. An isolated group uses νg = μg, giving zero contrast. Neighbor context thus retains appearance and relative spatial context in a 2560-dimensional vector.',
'A single fastPLS SIMPLS classifier with linear discriminant analysis, autoscaling and 24 components identifies the tissue group [6]. The fixed training partition supplies 967 whole-group examples and 25 represented tissue roles. Numeric computations use the CPU implementation with float64 inputs. The seed is 314159. The target is the hidden source identity that generated the coarse group, rather than a majority label recomputed from its final fragments.',
'For a new user-trained model, Bora first builds anonymous group descriptors without reading a reference. The user can then supply an aligned expert semantic TIFF or native-coordinate GeoJSON and map its tissue values to model class IDs. The dominant expert class is assigned to a whole group only when annotated coverage and purity meet declared thresholds; ambiguous groups are excluded. This mask/GeoJSON route is a supported training interface, whereas the reported 967 examples used the known source identities above.',
'Training and prediction artifacts record feature shape, represented labels, implementation revision and model settings. Because the binding does not expose portable fitted-state serialization, the safe artifact retains training arrays and refits once when loaded. Startup fitting is charged separately from warm prediction. Training features remain subject to the source data terms; a model artifact is not a license to redistribute them.'
]),
('Compact vector export and cleanup', [
'The published Bora coverage converter uses its accurate profile, with a 500-native-pixel component filter and two-pixel shared-border simplification. These settings precede, and are distinct from, the final 20-pixel cleanup. Per-class geometry is repaired and dissolved. Shared boundary arcs are fitted once in physical coordinates using straight-line simplification and then polygonized into consistent neighboring faces.',
'Fitting uses at most 3 µm tolerance, at least 97% per-class agreement with the preceding predicted geometry and at most 1.5% tissue-area change. These are predicted-shape safeguards, not agreement constraints against expert annotation. Polygon parts and enclosed holes strictly below 20 native pixels of area are removed or filled; exactly 20 is retained. Adjacent small parts are absorbed into surrounding tissue, while isolated specks on glass are removed to background. Continuous polygon area and native raster connectivity are both checked. The final raster is generated from the exported geometry, preventing disagreement between displayed polygons and the scored output.'
]),
('Experimental assessment', [
'The source registry contains 533 calibrated fields, partitioned into 349 training, 72 validation and 112 test images. Source identifiers link known duplicate images and TCGA cases across datasets. Other groups may represent slides, cores or ROIs rather than verified patients, and unresolved cross-source links remain a limitation. The tissue ontology retains 28 broad roles and their original provider labels. PUMA contributes to the training corpus but is excluded from both evaluation splits.',
'The multiclass evaluation comprises 116 eligible images: 42 BCSS, 26 TIGER, 24 GCSS and 24 IGNITE [7–10]. Their original split labels are 78 test and 38 validation. These previously inspected images constitute an exploratory evaluation, not independent final testing. Training targets are drawn only from the training partition. Evaluation roles absent from training are counted as errors, not removed from the endpoint. BCSS and TIGER share some source material; their results do not establish independent breast-cancer replication.',
'The registry has 326 training, 65 validation and 103 test source-group identifiers. The 116 eligible images represent 108 source groups; none of their known identifiers occurs in training, and the two evaluation subsets have no known identifier overlap. A secondary screen verified image SHA256s and compared 63-bit grayscale perceptual hashes for all 40,484 training–evaluation image pairs, flagging Hamming distance ≤4. These identifiers and global hashes may miss overlapping crops, rescans or slide-level relatedness, so the audit cannot certify patient independence. Source-group bootstrap intervals resample identifiers 10,000 times with a fixed seed and retain equal image weighting. They describe uncertainty within the reused development cohort.',
'A separate matched comparison assesses conditional geometry refinement on 39 eligible ROIs with identical coarse inputs, evidence and scoring masks. Comparators are unrefined tiles, soft-observation H&E graph-label diffusion and a convergence-qualified bilateral-solver reference notebook [3]. The graph control is not the original hard-seeded random walker. Bilateral settings use an 8-µm spatial scale, luminance/chroma scales of four, λ = 128 and relative conjugate-gradient tolerance 10⁻⁵; all retained class solves converge. The timings characterize this reference implementation, not the optimized implementation associated with the bilateral paper.',
'For matched whole-group controls, the same 967 training groups and 398 evaluation groups were paired with four comparisons: pooled group mean and spatial pooling with the native fastPLS classifier, neighbor contrast without group appearance using fastPLS, and the full Neighbor context descriptor with a standardized ridge classifier. The first two were carried through identical geometry and vector export; the latter two compare group identity only. The fastPLS reference was independently refitted from exported frozen descriptors and reproduced 350 correct group identities exactly. No target annotations were fitted in any classifier. These controls test the descriptor and classifier within this conditional pipeline; they do not represent contemporary dense segmentation systems.',
'Pixel macro intersection-over-union averages tissue roles with nonzero union on valid annotated pixels, including spurious predicted roles. Unannotated/unmapped pixels are ignored, but annotated tissue lost to foreground estimation remains error. Each image contributes equally to the image-mean endpoint. Whole-group accuracy and balanced accuracy respectively summarize correct group identities and mean recall across observed reference roles.',
'For tumor–stroma boundaries, four-neighbor pixel adjacency defines a direct interface between tumor and either stroma or inflamed stroma within annotated tissue. Precision and recall match interface pixels within 8 µm in physical coordinates; their harmonic mean is reported for images with an expert interface. A missing predicted interface scores zero. Average symmetric surface distance is calculated only where both interfaces exist. As a diagnostic of refinement geometry, we also score the coarse map after decoding its class identities with the expert-origin key; this privileged comparison is not a deployable classifier.',
'CPU measurements use an Intel Core i7-13700 with two numerical threads. Controlled refinement timing is the median of three repeats after warm-up and excludes encoder inference, RGB construction and loading. Cached pipeline time includes recorded appearance-field refinement, pooling, context construction, warm identification, vector export, fitting and final cleanup. Reused identical geometry retains its measured time charge. Cold loading, encoder extraction and startup fitting are excluded.',
'For an exploratory scale test, a public 10x Genomics colon-cancer whole-slide image [13] was processed in 228 non-overlapping 4096-pixel native cores with 512-pixel context. OME spacing was 0.273812 µm/pixel per axis and encoder input was resampled to approximately 0.5 µm/pixel. An annotation-independent slide-wide tile clustering supplied the coarse map. The previously trained Neighbor context head assigned tissue classes to groups within each bounded core. A supplied partial expert GeoJSON was excluded from inference and used only for scoring. The exported group readout was selected on this same slide; its accuracy is descriptive. Twenty-three homogeneous or undersized cores used the documented coarse-palette evidence fallback.'
,'A source-to-target TIGER ROI pilot used 69 expert-labelled, high-purity 45-µm tiles from one ROI to train a frozen fastPLS tile model. Predicted anonymous tiles on another ROI entered the unchanged RGB conversion and Bora refinement with calibrated Virchow2 tokens and H&E. Each final group received the area-majority predicted tile class, distinct from the main Neighbor context classifier. Target annotation was withheld until native mask and GeoJSON export. The selected two-role target was previously inspected; patient independence is unverified.'
]),
('Results', [
'In the matched 39-region comparison, Bora achieved mean native-mask macro IoU of 0.6628, versus 0.5862 for coarse tiles, 0.6061 for graph-label diffusion and 0.6435 for the bilateral reference (Table 1). Mean refinement time was 0.1676 seconds, compared with 0.0434 seconds for graph diffusion and 0.4569 seconds for the bilateral notebook. Equal source-group averaging gave a Bora-minus-bilateral IoU difference of 0.0208, with a 95% source-group bootstrap interval of 0.0093–0.0337. This interval describes the inspected cohort and is not a generalization guarantee.',
'On the 116-image assessment, Neighbor context correctly identified 350 of 398 whole groups (87.9%), with balanced accuracy of 68.3%. Mean final GeoJSON IoU was 0.5867; the preceding raster had mean IoU 0.5855. Dataset-specific final IoU ranged from 0.4803 on IGNITE to 0.7178 on GCSS (Table 2). Two observed reference roles lacked represented training examples. The difference between overall and balanced accuracy indicates that frequent tissue groups are identified more reliably than rare roles.',
'The originally designated test subset contained 78 images from 72 known source groups and had mean final IoU 0.5967 (95% source-group bootstrap interval 0.5411–0.6505). The validation subset contained 38 images from 36 groups and had mean 0.5661 (0.5034–0.6290). Both subsets were used during method development and cannot be interpreted as untouched tests. The whole-group pooled-mean and spatial-pooling controls correctly classified 340 and 341 of 398 groups, respectively, versus 350 for Neighbor context (Table 3). Their final IoUs were 0.5606 and 0.5596. The paired Neighbor-context differences were 0.0261 (−0.0002 to 0.0579) and 0.0271 (−0.0009 to 0.0593) under source-group resampling; these descriptive intervals include zero. Contrast-only fastPLS classified 341 groups and the non-PLS ridge control classified 335. Their group-only balanced accuracies were 0.573 and 0.562, versus 0.683 for Neighbor context. These two controls did not produce final masks.',
'The duplicate screen found no exact-image duplicate and no near-global-hash candidate at its prespecified threshold among the 40,484 training–evaluation pairs; it does not rule out partial overlap or shared patients.',
'Across annotated images, tumor and stroma had mean per-image IoUs of 0.767 and 0.704, respectively, but blood and blood vessels had means of 0.045 and 0.117. Four observed roles appeared in only one image, and two had no represented training examples. The overall annotated-unassigned fraction was 0.287% of reference pixels. In the lowest-IoU example, gcss_12, the classifier predicted neither the annotated tumor nor inflamed-stroma role; foreground loss was only 0.011% of reference pixels. This is primarily a tissue-identity failure rather than a background-removal failure. Full per-role support and IoU are supplied in the supplementary material.',
'Among 106 images with an expert tumor–stroma interface, mean 8-µm-tolerance boundary F1 was 0.571 (median 0.638); nine had no predicted interface. Expert-decoded coarse input had mean F1 0.354, yielding a paired final-minus-coarse difference of 0.216 (95% source-group bootstrap interval 0.175–0.256). This comparison tests refinement of an expert-derived partition and cannot establish improvement from predicted tile classes. Mean symmetric surface distance was 21.5 µm among images with both interfaces. Figure 3 displays one interface at native scale; the lowest-IoU case loses its interface through tissue misidentification.',
'The mean warm classifier time was 0.283 milliseconds per image, whereas mean recorded cached deployment was 2.071 seconds. The median was 1.139 seconds and the maximum 6.864 seconds. The classifier therefore contributes little to the measured image-level cost; millisecond identification does not imply a subsecond complete workflow.',
'Across the 116 final Neighbor context exports, compact fitting reduced boundary segments from 714281 to 326883, a 54.2% reduction. Relative to the preceding predicted geometry, mean minimum per-class agreement was 0.981 (minimum 0.970), with mean changed tissue area of 0.55% (maximum 1.42%). Polygon parts decreased from 4134 to 4079, so compactness involved some topology changes. Recorded final cleanup filled 14 small regions and 17 enclosed holes. Saved continuous-area and native-raster audits report no remaining subthreshold positive parts or enclosed holes. The 20-native-pixel cutoff spans approximately 1.25–20.46 µm² across evaluated spacings. These checks do not measure agreement with expert contours. Figures 1–3 and the public viewer expose representative and complete outputs.'
,'The whole-slide scale test covered 75,250 × 48,740 pixels (3.67 billion) and produced a native OME mask and 203-MB GeoJSON. Its partial reference supplied 298.5 million singly mapped pixels (8.14% of the slide). Macro Dice was 0.486 and macro IoU 0.375 across five mapped roles; normal gland had Dice 0 because the transferred head predicted none (Supplementary Figure S1). Independent rasterization of all 228 exported windows matched the final mask exactly, with no polygon part or enclosed hole below 20 native pixels. The encoder/refinement stage took 3005.8 s and a separate CPU vector pass took 1407.4 s, excluding earlier checkpointed work and other overhead. This single-slide, reference-selected result demonstrates bounded processing and a domain-transfer failure, not a generalization estimate or cold end-to-end throughput.'
,'In the source-to-target ROI, macro IoU on 1,402,860 annotated pixels increased from 0.603 for predicted tiles to 0.771 after Bora (difference 0.168; Supplementary Figure S2). Tumor and stroma IoUs were 0.791 and 0.751. The native raster matched the two-feature GeoJSON exactly; no part or hole was below 20 pixels. This selected ROI is an applicability example, not a generalization estimate.'
]),
('Discussion', [
'Bora offers a compact computational path from coarse tissue assignments to interpretable geometric regions. Its appearance-field model regularizes image evidence rather than directly smoothing the categorical label plane. All supplied groups can compete throughout supported tissue, which avoids choosing an arbitrary edit margin. Physical calibration connects encoder context, field spacing and boundary fitting across differently sampled datasets.',
'Neighbor context adds an explicit tissue relationship to whole-group appearance without a separate message-passing network. Weighting by actual interface length makes its neighborhood definition interpretable. The regional descriptor and the shared-boundary export are useful architectural choices, but graph regularization, partial least squares and polygon simplification are established methods. Bora integrates these components into a specified workflow for calibrated conditional tissue refinement and identification. The present evidence does not establish superiority over contemporary dense segmentation or feature-upsampling systems.',
'The matched controls suggest that combining group appearance with physical neighbor contrast may improve identity assignment. Contrast alone performed less well, and changing the estimator to ridge also reduced balanced accuracy. Yet only the first two controls have final segmentation outputs, and the paired final-IoU intervals include zero. The reused cohort further limits causal and generalization claims. Direct Virchow2 distillation [11] and multi-resolution promptable segmentation [12] address neighboring goals, but neither is an interchangeable comparator without harmonizing inputs, training supervision, tissue ontology and hardware.',
'The benchmark quantifies a controlled task. Anonymizing class identities does not remove expert information from group number, assignment or geometry, and the RGB representation uses those same coarse groups. Consequently, the results do not demonstrate autonomous segmentation, robustness to arbitrary upstream errors or recovery of a class absent from the input partition. Prior inspection of evaluation images, incomplete patient linkage, rare classes and the absence of some training roles further limit generalization claims.',
'The colon test demonstrates one annotation-independent clustering route but has a partial reference and same-slide-selected readout. The TIGER pilot passes source-trained fastPLS tile predictions through Bora without target annotations during inference, but uses selected two-role ROIs and an area-majority group readout. A convincing transfer study needs multiple source WSIs, patient-independent target WSIs and complete evaluation at another institution.',
'Compact polygons support inspection and downstream spatial analysis, but geometric compliance is not proof of preserved biological detail. The 20-pixel cutoff has different physical area at different native spacings, and some polygon parts are lost during fitting. Despite improved mean boundary F1 against expert-decoded coarse input, missing interfaces and a 21.5-µm mean distance among retained interfaces show that important morphology is not reliably preserved. Pathology review, topology-specific analysis and downstream morphometry remain needed. Full-slide latency, memory use and uncached encoder cost also require direct measurement. Clinical utility and independent external performance remain unestablished.'
' Per-role results reveal that the mean IoU conceals severe failures on vascular and rare roles. This further motivates predefined clinically important tissue interfaces and a genuinely new-patient test with an upstream predicted map.',
'The colon whole-slide test shows that processing scale and semantic transfer must be assessed separately. Windowed execution and exact native-coordinate export succeeded, but the classifier missed normal gland completely. Its partial annotation and same-slide readout selection preclude a claim of independent accuracy. The measured stages also omit total cold startup, image I/O and peak memory.',
]),
('Conclusions', [
'Bora combines calibrated appearance-field refinement, whole-group Neighbor context identification and compact shared-boundary export. Exploratory benchmarks show recovery of tissue structure from expert-derived coarse inputs, efficient warm classification and one completed predicted-tile ROI example. The public viewer exposes the assessed images and final vectors. Independent patient-level testing with predicted coarse maps remains necessary for broader deployment claims.'
]),
('Data and code availability', [
'Bora is written in Python with a bundled MIT-licensed fastPLS C++ CPU core. The application implementation at commit 370ba9df782b7b1858d0222f13aeeedd2ca8524b (https://github.com/tkcaccia/Bora) documents training a model from user-provided semantic masks or GeoJSON and applying it to separate tile clusters; this commit remains on a review branch at the time of writing. Acquisition, benchmark, scoring and figure-generation code are in the companion repository (https://github.com/tkcaccia/Bora-Analysis). Its machine-readable training-source manifest identifies all 349 original fields, native crops, class mappings and checksums, and its acquisition guide links provider releases. The interactive benchmark, whole-slide scale viewer and predicted-tile ROI example are available at https://tkcaccia.github.io/Bora-Analysis/. Dataset providers retain their image and annotation rights [7–10,13]; attribution and reuse terms accompany the viewer. Raw archives, large feature caches and gated encoder checkpoints are obtained separately under their original provider terms.'
]),
('Declaration of generative AI assistance', [
'OpenAI Codex was used to assist software development and manuscript drafting. Scientific claims, reported data and the decision to submit remain the responsibility of the human authors. The histology figures comprise existing dataset images and computational segmentation outputs; no generative image model was used to synthesize the illustrated tissue.'
])]

REFERENCES = [
'Zimmermann E, Vorontsov E, Viret J, et al. Virchow2 Scaling Self Supervised Mixed Magnification Models in Pathology. 2024. arXiv:2408.00738. https://arxiv.org/abs/2408.00738',
'Grady L. Random walks for image segmentation. IEEE Trans Pattern Anal Mach Intell. 2006;28:1768–1783. https://doi.org/10.1109/TPAMI.2006.233',
'Barron JT, Poole B. The Fast Bilateral Solver. ECCV. 2016. https://arxiv.org/abs/1511.03296',
'Wang L, He L, Mishra A, Li C. Active contours driven by local Gaussian distribution fitting energy. Signal Processing. 2009;89:2435–2447. https://doi.org/10.1016/j.sigpro.2009.03.014',
'Fu S, Hamilton M, Brandt L, Feldman A, Zhang Z, Freeman WT. FeatUp A Model Agnostic Framework for Features at Any Resolution. ICLR. 2024. https://arxiv.org/abs/2403.10516',
'Cacciatore S. fastPLS Python bindings and C++ numerical core. Version 0.3.0. MIT license. https://github.com/tkcaccia/fastPLS-py',
'Amgad M, Elfandy H, Hussein H, et al. Structured crowdsourcing enables convolutional segmentation of histology images. Bioinformatics. 2019;35:3461–3467. https://doi.org/10.1093/bioinformatics/btz083',
'van Rijthoven M, Aswolinskiy W, Tessier L, et al. TIGER training dataset ROI-level annotations of WSIROIS subset. Zenodo. 2022. https://doi.org/10.5281/zenodo.6014422',
'Shi J. Gastric Cancer Semantic Segmentation dataset accompanying Semi-Supervised Pixel Contrastive Learning Framework for Tissue Segmentation in Histopathological Image. Zenodo. 2023. https://doi.org/10.5281/zenodo.8176354',
'Spronck JMA, van Eekelen L, van Midden D, et al. IGNITE data toolkit a tissue and cell-level annotated H&E and PD-L1 histopathology image dataset in non-small cell lung cancer. Version 2. Zenodo. 2025. https://doi.org/10.5281/zenodo.17735903',
'Lv J, Zhu Y, Okada S, et al. Computationally Efficient Pathology Segmentation using Knowledge Distillation from Foundation Models. arXiv:2609.03947. 2026. https://arxiv.org/abs/2609.03947',
'Liu H, Yang H, van Diest PJ, Pluim JPW, Veta M. WSI-SAM: Multi-resolution Segment Anything Model for histopathology whole-slide images. Proc MICCAI Workshop on Computational Pathology. 2024;254:25–37. https://proceedings.mlr.press/v254/liu24a.html',
'10x Genomics. Visium HD Spatial Gene Expression Library, Human Colorectal Cancer (FFPE). Space Ranger 4.0.1 dataset, 2025. https://www.10xgenomics.com/datasets/visium-hd-cytassist-gene-expression-libraries-of-human-crc-v4']


def paragraph(doc, text, style=None):
    return doc.add_paragraph(text, style)


def configure(doc):
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
        s.page_width = Inches(8.5); s.page_height = Inches(11)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption','List Bullet']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        fonts=style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if key.endswith('Theme'):del fonts.attrib[key]
        fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
        style.font.size=Pt(12)
        style.paragraph_format.line_spacing=2
        style.paragraph_format.space_after=Pt(0)
    doc.styles['Title'].font.size=Pt(18)
    doc.styles['Title'].paragraph_format.line_spacing=1.15
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(12)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Caption'].paragraph_format.line_spacing=1.15
    doc.styles['Caption'].font.size=Pt(10)
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for section in doc.sections:
        f=section.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)


def equation(doc, text):
    p=paragraph(doc,'');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    m=OxmlElement('m:oMath')
    def run(parent,text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);parent.append(r)
    cursor=0
    for match in re.finditer(r'([A-Za-zμνℓ]+)_([A-Za-z]+)',text):
        run(m,text[cursor:match.start()]);sub=OxmlElement('m:sSub')
        base=OxmlElement('m:e');index=OxmlElement('m:sub');run(base,match[1]);run(index,match[2])
        sub.extend([base,index]);m.append(sub);cursor=match.end()
    run(m,text[cursor:]);p._p.append(m)


def table(doc, caption, headers, rows, widths, new_page=False):
    p=paragraph(doc,caption,'Caption');p.paragraph_format.keep_with_next=True
    p.paragraph_format.page_break_before=new_page
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,(c,w) in enumerate(zip(row.cells,widths)):
            c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            props.append(b)
            if not i:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E6EBEF');props.append(fill)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.15
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if not j else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(10);r.bold=not i
    paragraph(doc,'')


def figure(site,c,out):
    try:font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',25)
    except OSError:font=ImageFont.load_default(size=25)
    small=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',21)
    canvas=Image.new('RGB',(1800,1350),'white');draw=ImageDraw.Draw(canvas)
    panels=[('Original H&E','original'),('Expert annotation','expert'),('Anonymous input tiles','tiles'),('Virchow2 RGB evidence','converted'),('Refined anonymous groups','bora'),('Neighbor context result',None)]
    for i,(title,file) in enumerate(panels):
        x=600*(i%3);y=590*(i//3)
        draw.text((x+15,y+10),title,fill='black',font=font)
        path=site/c['asset']/(file+'.jpg') if file else site/c['models'][0]['preview']
        image=Image.open(path).convert('RGB');image.thumbnail((570,520),Image.Resampling.LANCZOS)
        ox=x+(600-image.width)//2;oy=y+52
        canvas.paste(image,(ox,oy))
        length=100/c['native_mpp_xy'][0]*image.width/c['native_shape_yx'][1]
        bx=ox+18;by=oy+image.height-40
        draw.rectangle((bx-6,by-29,bx+max(length,85)+8,by+9),fill='white')
        draw.line((bx,by,bx+length,by),fill='black',width=5)
        draw.text((bx,by-26),'100 µm',fill='black',font=small)
    roles={g['role'] for g in c['expert_legend'] if g['role']}
    vector=json.loads((site/c['models'][0]['vector']).read_text())
    roles|={int(f['properties']['value']) for f in vector['features']}
    data=json.loads((site/'data.json').read_text())
    xx=15;yy=1190
    for role in sorted(roles):
        text=data['roles'][role-1].replace('_',' ')
        width=draw.textlength(text,font=small)+45
        if xx+width>1785:xx=15;yy+=36
        draw.rectangle((xx,yy+5,xx+20,yy+25),fill=tuple(data['colours'][role]))
        draw.text((xx+28,yy),text,fill='black',font=small);xx+=width
    canvas.save(out,dpi=(300,300))


def build(site,author_block,output,reanalysis,classifier_controls,boundaries,vector_fidelity,duplicate_screen,boundary_figure):
    output.mkdir(parents=True,exist_ok=True);(output/'figures').mkdir(exist_ok=True)
    data=json.loads((site/'data.json').read_text())
    extra=json.loads(reanalysis.read_text())
    group_only=json.loads(classifier_controls.read_text())['results']
    boundary=json.loads(boundaries.read_text())['summary']
    vector=json.loads(vector_fidelity.read_text())['summary']
    duplicates=json.loads(duplicate_screen.read_text())['summary']
    assert data['cohort_cases']==116 and data['models'][0]['name']=='Neighbor context'
    assert extra['provenance']['images']==116
    assert abs(extra['comparators']['neighbor']['mean_final_iou']-data['models'][0]['mean_final_geojson_iou'])<1e-8
    assert extra['comparators']['neighbor']['group_correct']==350
    assert group_only['neighbor_fastpls']['correct_groups']==350
    assert group_only['contrast_only_fastpls']['correct_groups']==341
    assert group_only['neighbor_ridge']['correct_groups']==335
    assert boundary['cases_with_reference_interface']==106 and boundary['cases_with_missing_predicted_interface']==9
    assert vector['images']==116 and vector['fallbacks']==0
    assert duplicates['pairs_screened']==40484 and duplicates['exact_image_duplicate_pairs']==0 and duplicates['perceptual_hash_candidates_at_distance_4']==0
    doc=Document();configure(doc)
    doc.add_paragraph(TITLE,'Title')
    authors=json.loads(author_block.read_text())
    assert 'Giulia Canarutto⁴,¹¹' in authors['authors']
    assert '⁴ Computational Biology Group, International Centre for Genetic Engineering and Biotechnology, Trieste, Italy.' in authors['affiliations']
    assert '¹¹ Life Science Department (DSV), University of Trieste, Trieste, Italy.' in authors['affiliations']
    paragraph(doc,authors['authors'])
    paragraph(doc,authors['affiliations'])
    abstract_heading=doc.add_heading('Abstract',1)
    abstract_heading.paragraph_format.page_break_before=True
    for heading,text in ABSTRACT:
        p=paragraph(doc,'');p.add_run(heading+' ').bold=True;p.add_run(text)
    paragraph(doc,'Keywords: computational pathology; tissue segmentation; spatial context; partial least squares; physical calibration; GeoJSON')
    doc.add_page_break()
    subheads={'Inputs and physical calibration','Virchow2 evidence conversion','H and E guided appearance fields','Neighbor context tissue identification','Compact vector export and cleanup'}
    body=[]
    for heading,paragraphs in SECTIONS:
        doc.add_heading(heading,2 if heading in subheads else 1)
        for i,text in enumerate(paragraphs):
            paragraph(doc,text);body.append(text)
            if heading=='H and E guided appearance fields' and i==1:equation(doc,'(D_k + Lw + ρ I) F_k = B_k + ρ 1 μ_kᵀ')
            if heading=='Neighbor context tissue identification' and i==1:equation(doc,'x_g = [ μ_g , μ_g − ν_g ]')
    doc.add_heading('References',1)
    for i,text in enumerate(REFERENCES,1):paragraph(doc,f'[{i}] {text}')
    table(doc,'Table 1 Matched conditional refinement on 39 regions. Times exclude encoder inference, evidence construction and loading.',
          ['Method','Mean macro IoU','Mean refinement s'],[
              ['Coarse input tiles','0.5862','0.0010'],
              ['H&E graph-label diffusion','0.6061','0.0434'],
              ['Bilateral reference notebook','0.6435','0.4569'],
              ['Bora appearance fields','0.6628','0.1676']],[3.3,1.5,1.7],new_page=True)
    rows=[]
    for ds in ['bcss','tiger','gcss','ignite']:
        cs=[c for c in data['cases'] if c['dataset']==ds]
        mpp=[v for c in cs for v in c['native_mpp_xy']]
        resolution=f'{min(mpp):.3f}' if max(mpp)==min(mpp) else f'{min(mpp):.3f}–{max(mpp):.3f}'
        rows.append([ds.upper(),len(cs),resolution,f"{data['models'][0]['per_dataset'][ds]:.4f}"])
    table(doc,'Table 2 Neighbor context evaluation. All fields are ROIs or WSI-derived windows, not complete-slide performance measurements. Native per-axis spacing is available for every image in the viewer.',
          ['Dataset','Images','Native µm per pixel','Mean final IoU'],rows,[1.4,1,2.2,1.9])
    controls=extra['comparators']
    table(doc,'Table 3 Matched whole-group controls on the 116 development images. Accuracy uses 398 groups. The first three rows include final mask export; the last two compare classification only. Cached deployment excludes encoder extraction.',
          ['Descriptor','Correct groups','Balanced accuracy','Final macro IoU','Cached s/image'],
          [[label,f"{controls[key]['group_correct']}/398",f"{controls[key]['balanced_group_accuracy']:.3f}",
            f"{controls[key]['mean_final_iou']:.3f}",f"{controls[key]['mean_cached_deployment_seconds']:.3f}"]
           for label,key in [('Pooled mean','mean'),('Spatial pooling','spatial'),('Neighbor context','neighbor')]]+
          [[label,f"{group_only[key]['correct_groups']}/398",f"{group_only[key]['balanced_accuracy']:.3f}",'—','—']
           for label,key in [('Contrast only','contrast_only_fastpls'),('Neighbor ridge','neighbor_ridge')]],
          [1.65,1.15,1.35,1.25,1.1])
    ordered=sorted(data['cases'],key=lambda c:c['models'][0]['iou'])
    examples=[ordered[len(ordered)//2],ordered[0]]
    for i,c in enumerate(examples,1):
        path=output/'figures'/f'figure-{i}.png';figure(site,c,path)
        doc.add_page_break()
        p=paragraph(doc,'');p.paragraph_format.keep_with_next=True
        run=p.add_run();picture=run.add_picture(str(path),width=Inches(6.5))
        picture._inline.docPr.set('descr',
            f"Six-panel histology comparison for {c['case']}: original H and E, expert annotation, anonymous tile input, computational RGB evidence, refined anonymous groups and classified output. "
            + ("This is the median-scoring example." if i==1 else "The low-scoring example loses tumor and inflamed-stroma identities."))
        caption=f"Figure {i} {'Median-ranked' if i==1 else 'Lowest-IoU'} example {c['case']} with final macro IoU {c['models'][0]['iou']:.3f}. Expert annotations, input tiles and final masks are overlaid on H&E. Each scale bar denotes 100 µm. Anonymous colors identify groups; the tissue legend applies to expert and classified masks. RGB evidence is a computational representation, not measured molecular signal. {'The low-scoring image mainly illustrates class confusion, with no predicted tumor or inflamed-stroma role.' if i==2 else 'The public viewer includes all 116 images.'}"
        paragraph(doc,caption,'Caption')
    doc.add_page_break()
    p=paragraph(doc,'');p.paragraph_format.keep_with_next=True
    picture=p.add_run().add_picture(str(boundary_figure),width=Inches(6.5))
    picture._inline.docPr.set('descr','Four-panel native-scale tumor–stroma interface comparison on the same H&E crop: original, expert boundary, expert-decoded coarse boundary, and Bora boundary, each with a 50-micrometre scale bar.')
    paragraph(doc,'Figure 3 Native-scale tumor–stroma boundary example from IGNITE. The expert contour is yellow, expert-decoded coarse contour cyan and Bora final contour magenta. The coarse map is decoded using the expert-origin identity key solely for this diagnostic; it is not a deployable tile classifier. Each panel shows the same tissue crop and a 50 µm scale bar. Across the 106 images with an expert interface, mean boundary F1 at 8 µm tolerance was 0.571 for Bora and 0.354 for expert-decoded coarse input.','Caption')
    doc.core_properties.title=TITLE;doc.core_properties.author='Bora research team'
    doc.core_properties.subject='Conditional tissue refinement and Neighbor context identification'
    path=output/'Bora_CMPB_manuscript.docx';doc.save(path)
    highlights=Document();configure(highlights);highlights.add_paragraph('Highlights','Title')
    bullets=['Bora refines coarse tissue maps using physically calibrated appearance fields.',
             'Neighbor context identifies complete tissue groups with one fastPLS head.',
             'Shared-boundary fitting produces compact polygons in native coordinates.',
             'A public viewer exposes all 116 assessed images, masks and scores.']
    assert all(len(x)<=85 for x in bullets)
    for text in bullets:paragraph(highlights,text,'List Bullet')
    highlights.save(output/'Bora_CMPB_highlights.docx')
    words=len(re.findall(r'\b\S+\b',' '.join(body)))
    abstract_words=len(' '.join(h+' '+t for h,t in ABSTRACT).split())
    assert abstract_words<=250
    assert words<=3800
    all_text='\n'.join(p.text for p in doc.paragraphs)
    assert not re.search(r'\bcycle(?:s)?\b|\bevolution\b|\bevolve\w*\b|\bP\d{2}\b|\bR\d{2}\b|\bC\d{2}-\d{2}\b',all_text,re.I)
    (output/'MANUSCRIPT_CHECK.json').write_text(json.dumps(dict(body_words=words,abstract_words=abstract_words,cohort=116,groups=398,correct_groups=350,examples=[c['case'] for c in examples],structured_abstract=True,internal_method_codes=False,development_history_in_manuscript=False,independent_validation_claimed=False,whole_slide_scale_test_reported=True,exploratory_roi_transfer_reported=True,wsi_transfer_results_claimed=False),indent=2)+'\n')
    (output/'AUTHOR_ACTIONS.txt').write_text('Before submission, the authors must confirm authorship and affiliations, corresponding author contact, funding, competing interests, CRediT contributions, institutional determination for public-data reuse, and the final AI-assistance declaration. No statement of completed author approval or an ethics approval identifier has been invented. The article reports exploratory previously inspected data; independent external assessment and a matched upstream-prediction study remain the principal scientific limitations. This document has not been submitted.\n')
    print(json.dumps(dict(manuscript=str(path),body_words=words,abstract_words=abstract_words)))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--site',type=Path,required=True)
    parser.add_argument('--author-block',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--reanalysis',type=Path,required=True)
    parser.add_argument('--classifier-controls',type=Path,required=True)
    parser.add_argument('--boundaries',type=Path,required=True)
    parser.add_argument('--vector-fidelity',type=Path,required=True)
    parser.add_argument('--duplicate-screen',type=Path,required=True)
    parser.add_argument('--boundary-figure',type=Path,required=True)
    args=parser.parse_args();build(args.site.resolve(),args.author_block.resolve(),args.output.resolve(),args.reanalysis.resolve(),args.classifier_controls.resolve(),args.boundaries.resolve(),args.vector_fidelity.resolve(),args.duplicate_screen.resolve(),args.boundary_figure.resolve())
