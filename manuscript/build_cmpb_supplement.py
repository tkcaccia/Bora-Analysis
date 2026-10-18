"""Build the descriptive statistical supplement for the CMPB article."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from build_cmpb_manuscript import configure, paragraph, table


def build(reanalysis: Path, classifier_controls: Path, boundaries: Path, vector_fidelity: Path, output: Path) -> None:
    evidence = json.loads(reanalysis.read_text())
    group_only = json.loads(classifier_controls.read_text())['results']
    boundary = json.loads(boundaries.read_text())['summary']
    vector = json.loads(vector_fidelity.read_text())['summary']
    doc = Document(); configure(doc)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.82)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    doc.add_paragraph('Bora supplementary methods and exploratory results', 'Title')
    paragraph(doc, 'This supplement documents the fixed cohort and descriptive reanalysis accompanying the article on conditional tissue refinement. All figures come from previously inspected development images; intervals describe variation within that cohort and do not validate generalization to new patients.')
    doc.add_heading('Cohort and statistical unit', 1)
    paragraph(doc, 'The registry contains 533 calibrated fields: 349 train, 72 validation and 112 test. Eligibility and the exclusion of PUMA from evaluation leave 116 images across BCSS, TIGER, GCSS and IGNITE. The 116 images map to 108 known source-group identifiers. No identifier is shared with training, and validation and test identifiers do not overlap. Provider linkage is incomplete, so an identifier is not always a verified patient. The original split labels are retained for descriptive reporting, although both subsets were inspected during method development.')
    paragraph(doc, 'For each stratum we report the equal-image mean of per-image macro IoU. A 95% interval was formed from 10,000 bootstrap resamples of source-group identifiers with replacement (seed 19017), carrying all images within each selected identifier. These are descriptive intervals, and cross-provider patient overlap may remain unresolved. Unassigned-reference fraction divides annotated pixels mapped to final label zero by all valid annotated pixels.')
    c=evidence['cohort']
    table(doc, 'Table S1 Cohort-specific final mask scores and source-group bootstrap intervals.',
          ['Stratum', 'Images', 'Source groups', 'Mean macro IoU', '95% interval'],
          [[key.upper() if key in ('bcss','tiger','gcss','ignite') else key.title(),
            c[key]['images'], c[key]['source_groups'], f"{c[key]['mean_final_iou']:.3f}",
            f"{c[key]['source_group_bootstrap_95'][0]:.3f}–{c[key]['source_group_bootstrap_95'][1]:.3f}"]
           for key in ('all','test','validation','bcss','tiger','gcss','ignite')],
          [1.4,0.75,1.2,1.4,1.75])
    doc.add_heading('Matched whole-group controls', 1)
    paragraph(doc, 'All controls use the same frozen 967 training groups, tissue ontology and 398 evaluated groups on 116 images. Pooled mean uses the group-average Virchow2 representation. Spatial pooling aggregates spatial tissue evidence. Neighbor context concatenates the group mean and its contrast to neighbors weighted by physical shared-boundary length. Contrast only omits group appearance. The ridge control replaces fastPLS with a standardized linear classifier on the full Neighbor context descriptor. Pooled mean and spatial pooling used the same final geometry and vector export; contrast only and ridge were evaluated for group identity only. The native fastPLS reference was independently refitted from the frozen matrix and reproduced 350 correct identities.')
    k=evidence['comparators']
    table(doc, 'Table S2 Matched classifier controls on the 398 whole-group targets.',
          ['Descriptor','Correct groups','Balanced accuracy','Final mean IoU','Cached s/image'],
          [[label,f"{k[key]['group_correct']}/398",f"{k[key]['balanced_group_accuracy']:.3f}",
            f"{k[key]['mean_final_iou']:.3f}",f"{k[key]['mean_cached_deployment_seconds']:.3f}"]
           for label,key in [('Pooled mean','mean'),('Spatial pooling','spatial'),('Neighbor context','neighbor')]]+
          [[label,f"{group_only[key]['correct_groups']}/398",f"{group_only[key]['balanced_accuracy']:.3f}",'—','—']
           for label,key in [('Contrast only','contrast_only_fastpls'),('Neighbor ridge','neighbor_ridge')]],
          [1.65,1.15,1.35,1.25,1.1])
    paragraph(doc, 'The paired mean final-IoU difference for Neighbor context versus pooled mean was 0.0261 (95% source-group bootstrap interval −0.0002 to 0.0579). Versus spatial pooling, it was 0.0271 (−0.0009 to 0.0593). Both intervals include zero. These controls support a possible benefit of neighborhood contrast but do not establish an independent performance gain. Encoder extraction and evidence creation were excluded from cached timings.')
    doc.add_heading('Tissue-role support and failure analysis', 1)
    paragraph(doc, 'For a tissue role, the denominator is the number of images with at least one expert-annotated pixel of that role. Per-role IoU is averaged equally across those images. Values for sparsely represented roles are unstable; they should not be interpreted as population estimates. The ontology maps provider labels to broad roles and retains the original provider name in the benchmark registry. A role absent from training remains in evaluation and counts as an error.')
    roles=evidence['per_role']
    table(doc, 'Table S3 Per-role support and IoU for the final Neighbor context polygons.',
          ['Role','Annotated images','Reference pixels','Mean IoU'],
          [[roles[key]['name'].replace('_',' '), roles[key]['annotated_images'],
            f"{roles[key]['total_truth_pixels']:,}", f"{roles[key]['mean_iou_on_annotated_images']:.3f}"]
           for key in sorted(roles,key=lambda v:int(v))],
          [2.8,1.15,1.45,1.1])
    paragraph(doc, 'The lowest-scoring example, gcss_12, had final macro IoU 0.106. Tumor (226,403 annotated pixels) and inflamed stroma (229,249 pixels) received no predicted pixels, while the stroma role absorbed 412,551 predicted pixels and an inflammatory-infiltrate role appeared where it was not annotated. Only 72 of 636,294 valid reference pixels were left unassigned. The dominant failure is tissue identity, not foreground removal. The viewer shows the corresponding original image, expert labels, coarse input, converted representation and final mask.')
    doc.add_heading('Tumor–stroma interface and vector fidelity', 1)
    paragraph(doc, f"A direct four-neighbor tumor versus stroma/inflamed-stroma interface was present in {boundary['cases_with_reference_interface']} expert masks. With an 8-µm physical tolerance, final-mask mean F1 was {boundary['mean_f1_8um']:.3f} (median {boundary['median_f1_8um']:.3f}); {boundary['cases_with_missing_predicted_interface']} images had no predicted interface and received F1 zero. Mean symmetric surface distance was {boundary['mean_assd_um_when_both_present']:.1f} µm only where both interfaces existed. Decoding coarse-tile identity from the expert-origin key gave mean F1 {boundary['expert_decoded_coarse_mean_f1_8um']:.3f}; final-minus-coarse paired difference was {boundary['final_minus_expert_decoded_coarse_mean_f1_8um']:.3f} (95% source-group bootstrap interval {boundary['final_minus_coarse_source_group_bootstrap_95'][0]:.3f}–{boundary['final_minus_coarse_source_group_bootstrap_95'][1]:.3f}). The coarse comparison is a privileged geometry diagnostic, not an independent predictor.")
    paragraph(doc, f"Against each image's own pre-fitting predicted geometry, shared-boundary fitting had mean minimum per-class IoU {vector['mean_minimum_class_iou_vs_prefit_prediction']:.3f} (minimum {vector['minimum_class_iou_vs_prefit_prediction']:.3f}) and mean changed tissue area {100*vector['mean_changed_tissue_area_fraction']:.2f}% (maximum {100*vector['maximum_changed_tissue_area_fraction']:.2f}%). Polygon parts changed from {vector['parts_before']} to {vector['parts_after']}, with loss in {vector['images_with_lost_parts']} images; no export fell back to the prior geometry. These safeguards compare predicted shapes, not expert contours. Twenty native pixels correspond to approximately 1.25–20.46 µm² across the evaluated image spacings.")
    doc.add_heading('Reproducibility and unresolved measurements', 1)
    paragraph(doc, 'The public Bora-Analysis repository supplies the exact 349-field training-source manifest, acquisition instructions, ontology and processing code. The Bora application supplies a small-data tutorial for creating a model from an expert semantic TIFF or GeoJSON and applying it to another image. Large image archives and gated encoder weights remain with their original providers. The compact public viewer exposes all 116 evaluated images and vectors.')
    paragraph(doc, 'A patient-independent donor-to-target experiment with predicted or clustered coarse inputs, final masks for the contrast-only and non-PLS controls, expert-reviewed topology and downstream morphometry, and cold end-to-end encoder runtime and peak memory remain unmeasured. These measurements are needed before claims of generalization or whole-slide efficiency are made.')
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = 'Bora supplementary methods and exploratory results'
    doc.core_properties.author = 'Bora research team'
    doc.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--reanalysis', type=Path, required=True)
    parser.add_argument('--classifier-controls', type=Path, required=True)
    parser.add_argument('--boundaries', type=Path, required=True)
    parser.add_argument('--vector-fidelity', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.reanalysis, args.classifier_controls, args.boundaries, args.vector_fidelity, args.output)
