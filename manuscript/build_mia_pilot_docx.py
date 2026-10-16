"""Build the Medical Image Analysis pilot DOCX from its editable Markdown source."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from benchmark.runtime import BORA_ROOT


HERE = BORA_ROOT / 'docs'
SOURCE = HERE / 'Bora_MIA_pilot_manuscript.md'
TARGET = HERE / 'Bora_MIA_pilot_manuscript.docx'
FIGURES = {
    1: (HERE / 'figures/bora_workflow.png', 'Landscape Bora workflow with separate prospective tile-feature and evaluated benchmark branches, complementary boundary operations, and raster and GeoJSON outputs.'),
    2: (Path('/mnt/sata_ssd/Bora-benchmark/results/bora-pilot-independent-2026-10-16/pilot-macro-iou.png'), 'Macro intersection over union across tile widths and error rates in three pilot ROIs.'),
    3: (Path('/mnt/sata_ssd/Bora-benchmark/results/bora-pilot-independent-2026-10-16/pilot-class-iou.png'), 'Class intersection over union for BCSS, TIGER and WSSS4LUAD pilot ROIs.'),
    4: (HERE / 'figures/multi_tissue_mask_comparison.png', 'Three tissue datasets in rows, showing aligned reference, tile and Bora masks as semi-transparent colour fills over H&E.'),
    5: (Path('/mnt/sata_ssd/Bora-benchmark/results/bora-wsss-validation-2026-10-16/reviewer_cohort/delta_distribution.png'), 'ROI level distributions of macro IoU and class pair boundary F1 changes in WSSS4LUAD.'),
    6: (HERE / 'figures/breast_extension_changes.png', 'Source-group distributions of macro IoU and internal boundary F1 changes in the new BCSS and TIGER extensions.'),
}


def inline(paragraph, content: str) -> None:
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', content):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else
                                part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size in (('Title', 16), ('Heading 1', 12), ('Heading 2', 11)):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
    # Word's built-in Title style can carry a blue bottom border.
    title_ppr = doc.styles['Title'].element.pPr
    if title_ppr is not None:
        for border in title_ppr.findall(qn('w:pBdr')):
            title_ppr.remove(border)

    lines = SOURCE.read_text().splitlines()
    skip_pending = False
    table_rows = []

    def flush_table() -> None:
        if not table_rows:
            return
        table = doc.add_table(rows=1, cols=len(table_rows[0]))
        table.style = None
        table.autofit = False
        column_count = len(table_rows[0])
        if column_count == 2:
            widths = (1.15, 5.65)
        elif column_count == 5:
            widths = (1.10, 1.85, 2.25, .65, .95)
        elif column_count == 8:
            widths = ((1.10, .64, .60, .82, .82, .57, 1.10, 1.15)
                      if table_rows[0][0] == 'Source' else
                      (.60, .48, 1.05, 1.08, .48, 1.05, 1.08, .98))
        else:
            widths = None
        if widths:
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
        borders = OxmlElement('w:tblBorders')
        for edge, value, size in (('top', 'single', '5'),
                                  ('bottom', 'double', '7'),
                                  ('left', 'nil', '0'),
                                  ('right', 'nil', '0'),
                                  ('insideH', 'nil', '0'),
                                  ('insideV', 'nil', '0')):
            element = OxmlElement(f'w:{edge}')
            element.set(qn('w:val'), value)
            if value != 'nil':
                element.set(qn('w:sz'), size)
                element.set(qn('w:color'), '666666')
            borders.append(element)
        table._tbl.tblPr.append(borders)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i, value in enumerate(table_rows[0]):
            cell = table.rows[0].cells[i]
            if widths:
                cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if (column_count == 5 and i >= 3) or column_count == 8:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, value)
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)
            cell_borders = OxmlElement('w:tcBorders')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '5')
            bottom.set(qn('w:color'), '666666')
            cell_borders.append(bottom)
            cell._tc.get_or_add_tcPr().append(cell_borders)
        for values in table_rows[1:]:
            cells = table.add_row().cells
            for i, value in enumerate(values):
                if widths:
                    cells[i].width = Inches(widths[i])
                cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cells[i].paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                if (column_count == 5 and i >= 3) or column_count == 8:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                inline(paragraph, value)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        doc.add_paragraph()
        table_rows.clear()

    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('|'):
            values = [cell.strip() for cell in line.strip('|').split('|')]
            if not all(set(v) <= set('-: ') for v in values):
                table_rows.append(values)
            continue
        flush_table()
        if line.startswith('## Submission items pending author input'):
            skip_pending = True
        if skip_pending or line.startswith('**Article draft') or line.startswith('**Affiliation verification:'):
            continue
        if line.startswith('**Authors, in requested order:**'):
            line = line.replace('Authors, in requested order:', 'Authors:')
        if line.startswith('# '):
            paragraph = doc.add_paragraph(style='Title')
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, line[2:])
        elif line.startswith('## '):
            inline(doc.add_paragraph(style='Heading 1'), line[3:])
        elif line.startswith('### '):
            inline(doc.add_paragraph(style='Heading 2'), line[4:])
        elif re.match(r'\*\*Figure [1-6]\.\*\*', line):
            number = int(re.match(r'\*\*Figure ([1-6])\.', line).group(1))
            if number == 4:
                doc.add_page_break()
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.keep_with_next = True
            path, description = FIGURES[number]
            graphic = picture.add_run().add_picture(str(path), width=Inches(6.7))
            graphic._inline.docPr.set('descr', description)
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(11)
            paragraph.add_run(f'Figure {number}.').bold = True
            paragraph.add_run(line.split('**', 2)[-1].split(' Figure file:')[0])
            for run in paragraph.runs:
                run.font.size = Pt(9)
        else:
            paragraph = doc.add_paragraph()
            if re.match(r'\*\*Table \d+\.\*\*', line):
                paragraph.paragraph_format.keep_with_next = True
            if line.startswith('The accurate GeoJSON configuration proved'):
                paragraph.paragraph_format.keep_together = True
            inline(paragraph, line)
    flush_table()
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Bora pilot manuscript • Medical Image Analysis draft')
    doc.save(TARGET)
    print(TARGET)


if __name__ == '__main__':
    main()
