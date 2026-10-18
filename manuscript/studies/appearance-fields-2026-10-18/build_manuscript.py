"""Build the versioned article draft from frozen evidence with editable math."""
from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
REPORTS=ROOT.parents[1]/'results'
SEG=REPORTS/'bora-segment-identification-2026-10-18-v2'
OLD=REPORTS/'bora-supervised-2026-10-18-v1'
load=lambda p:json.loads(p.read_text())
summary=load(SEG/'REPORT.json');registry=load(OLD/'REGISTRY.json');controlled=load(OLD/'CONTROLLED_FINAL.json')
val={m['family']:m for m in summary['results']['validation']['models']}
test={m['family']:m for m in summary['results']['followup_test']['models']}
priors={s:load(SEG/'frequency-control-v1'/s/'RESULTS.json')['ranking'][0] for s in ('validation','test')}
af={r['name']:r for r in controlled['ranking']}
doc=Document();section=doc.sections[0];section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=Inches(.72);section.bottom_margin=Inches(.70);section.left_margin=Inches(.79);section.right_margin=Inches(.79)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'Times New Roman')
    style._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].paragraph_format.space_after=Pt(10)
for name,size in [('Heading 1',13),('Heading 2',11.5)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(11);s.paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=False
header=section.header.paragraphs[0];header.text='Bora appearance fields  |  Article draft 18 October 2026';header.runs[0].font.size=Pt(8)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,level=1)
def hh(text):doc.add_heading(text,level=2)
def equation(text):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath');run=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;run.append(t);math.append(run);para._p.append(math)
    para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(7)
def table(caption,headers,rows,widths):
    cap=p(caption,'Caption');cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for cell,label,w in zip(t.rows[0].cells,headers,widths):cell.text=label;cell.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,value,w in zip(cells,row,widths):cell.text=str(value);cell.width=Inches(w)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders);margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF0');props.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(9.2);run.bold=i==0
    p('')
    return t
def figure(name,caption,width=6.55):
    para=doc.add_paragraph();para.paragraph_format.keep_with_next=True;para.add_run().add_picture(str(ROOT/'figures-corrected-segments'/name),width=Inches(width))
    p(caption,'Caption')
def reference(number,text,url):
    para=p(f'[{number}] {text} ');para.paragraph_format.space_after=Pt(5)
    for run in para.runs:run.font.size=Pt(9.5)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');props.append(sz);run.append(props)
    text=OxmlElement('w:t');text.text='Source';run.append(text);link.append(run);para._p.append(link)

p('Bora appearance fields for tissue segmentation and segment identification','Title')
p('Martin Ocharo¹,³; Dalia Ahmed¹,²; Moussa Kassim Mohammed¹,³; Awonke Fuzile³; Aluwani Nemutavhani³; Enya Steyn³; Dupe Ojo¹; Silvano Piazza⁴; Dinesh Gupta⁵; Stefano Cacciatore¹')
aff=p('¹ Bioinformatics Unit, ICGEB, Cape Town, South Africa. ² Histopathology and Cytology Department, Faculty of Medical Laboratory Science, Omdurman Ahlia University, Sudan. ³ University of Cape Town, Cape Town, South Africa. ⁴ Computational Biology Group, ICGEB, Trieste, Italy. ⁵ Translational Bioinformatics Group, ICGEB, New Delhi, India.')
for run in aff.runs:run.font.size=Pt(9.5)
p('Working article draft. Author order and affiliations are retained from the supplied manuscript. This version distinguishes controlled refinement from exploratory identification of whole tissue segments.','Caption')
hh('Abstract')
p('Coarse tissue maps derived from pathology tiles quantize tissue interfaces and leave anonymous groups without validated semantic identities. We develop Bora-AF, H&E-guided Appearance Fields, to refine supplied groups by fitting spatial class appearances and applying whole-field competition. A variance-adaptive extension, Bora-VAF, is evaluated separately. On 39 controlled development regions of interest, Bora-AF achieved mean native-pixel macro intersection over union (IoU) of 0.663 at 0.168 seconds per region, compared with 0.643 at 0.457 seconds for a convergence-qualified bilateral solver reference implementation. These timings exclude encoder inference and RGB construction. A separate annotation-free grouping route generated 67,008 connected Bora regions across 469 calibrated images. We then froze their boundaries and trained native fastPLS CPU models to identify each whole region. The validation-selected pooled-region classifier achieved 70.0% top-one identification and 96.7% top-five containment on 6,379 scored segments from 62 exploratory follow-up images; balanced accuracy was 49.0% and native semantic-mask IoU was 0.232. These results support efficient conditional refinement and feasible region-level identification, while exposing limits from mixed regions, class imbalance and ontology differences. The follow-up cohort had been inspected previously and is not an independent final test.')
p('Keywords: computational pathology; tissue segmentation; image-guided refinement; regional tissue identification; foundation model; partial least squares','Caption')
h('1 Introduction')
p('Pathology foundation models provide representations that can distinguish tissue architecture, but projecting one group assignment per tile onto an image gives a coarse geometric partition. Boundary refinement and tissue identification are separate problems. A refined region may still have an anonymous group ID; conversely, an accurate tile label does not establish the boundary or identity of a complete anatomical region. Virchow2 provides a useful feature source [1], but an embedding cluster is not automatically a named tissue compartment.')
p('Bora-AF treats refinement as the estimation of a spatial appearance model for every supplied group. The original H&E image guides how that appearance varies across space, while a fixed evidence RGB supplies the quantities being modeled. All groups compete over the tissue field. Reliable interiors influence the fit continuously rather than acting as immovable seeds or limiting changes to a prescribed boundary band. Bora-VAF provides an optional extension with fitted local variances and a different spatial regularizer.')
p('We also investigate the intended semantic output: one identity and a top-five shortlist for each connected tissue segment after Bora. Tiles provide the coarse grouping and may contribute votes inside a refined region. The final classifier operates on a complete frozen region, either through aggregated tile evidence or directly pooled features. Region identity is stored separately from tissue identity so neighboring segments assigned the same tissue type remain distinguishable.')
p('The study separates two questions. A controlled experiment measures recovery from reference-derived tile quantization using matched inputs and timings. An exploratory supervised experiment assesses the identification of annotation-free Bora regions. Their scores have different label spaces and starting conditions and are not pooled. This distinction is necessary to avoid interpreting successful conditional refinement as evidence of autonomous semantic segmentation.')
h('2 Related work and contribution')
p('Graph-based image regularization is established. Random walkers infer labels from seeded graphs [2], and the Fast Bilateral Solver optimizes confidence-weighted image-guided estimates [3]. Local Gaussian distribution fitting has already used spatially varying means and variances for segmentation [4]. Bora therefore does not claim novelty for quadratic regularization, local Gaussian models, or training-free refinement alone. Its methodological focus is to regularize conditional class appearances, then use the fitted appearances for categorical competition on physically calibrated pathology fields. The soft graph-label diffusion baseline tests the alternative of regularizing labels directly.')
p('Feature upsampling methods such as FeatUp, Upsample Anything and AnyUp address the spatial resolution of learned representations [5–7]. SegmentAnyTissue combines foundation-model representations and segmentation prompts [8]. These are relevant alternatives or components, but they were not executed in the present matched comparison. Quantitative superiority over them is not established.')
p('Lv et al. train Virchow2 segmentation teachers and distill their responses and features into smaller dense predictors [9]. Their broad objective of efficient pathology segmentation overlaps with ours. The deployment mechanisms differ: the distilled student is a learned pixel predictor and need not execute Virchow2 for each image, whereas our present route requires encoder features or their cache, refines an explicit coarse partition, and subsequently identifies whole connected regions. A matched independent evaluation against a distilled student remains necessary. The current fastPLS head is a supervised regional identifier, not a distilled dense decoder [10].')
h('3 Methods')
hh('3.1 Physical sampling and input evidence')
p('Inputs are an aligned H&E image, a categorical coarse map, a tissue support estimate and explicit per-axis pixel spacing. All supervised feature inputs use full, unselected 1,280-dimensional Virchow2 spatial-token grids from a fixed encoder revision. At the 0.5 µm per pixel working grid, 14-pixel token spacing corresponds to 7 µm and a 224-pixel encoder window covers 112 µm. These are sampling and context sizes, not measurements of optical resolving power. Native and working grids are recorded separately; interpolation cannot recover missing optical detail. Unknown calibration, small fields and insufficient physical tile counts exclude a source from this physically calibrated study.')
p('In the autonomous region experiment, features are pooled in 45 µm physical fields. Within each image, PCA retains at most 16 nondegenerate directions, whitens them, and supplies a fixed four-cluster KMeans model with 10 initializations and seed 314159. Fitting fields require at least 50% H&E tissue and at least 36 µm extent along each axis. Every clipped edge tile is still predicted. This prevents tiny clipped corners from defining spurious full-field groups. Four groups are a fixed proof-of-concept setting and are not assumed to equal the true number of tissue types.')
p('We retain the existing predictable latent structure RGB conversion, algorithm 19. It selects 64 informative feature dimensions relative to the anonymous coarse groups. A ridge model predicts their values from neighbors at token offsets 1, 2 and 4 using four spatial folds and a four-token exclusion buffer. Group probabilities from the completed features are rendered with fixed colors. Encoder receptive fields overlap despite this buffer. The output is an evidence projection, not a reconstruction of H&E or an independent molecular measurement. The conversion formula is unchanged. No expert reference, semantic vocabulary or earlier reference-conditioned feature-selection cache enters autonomous geometry generation.')
figure('figure1-region-workflow.png','Figure 1. Geometry precedes identity. Anonymous physical-tile groups and fixed RGB evidence enter Bora-AF. All connected native region IDs and evidence are frozen before semantic training labels are attached. Classification returns one ranked tissue list for each whole region. This diagram depicts the autonomous region experiment; the separate controlled experiment supplies reference-derived coarse groups.')
hh('3.2 Bora AF conditional appearance fields')
p('We name J01 Bora-AF, H&E-guided Appearance Fields. Evidence is analyzed on an approximately 1 µm grid, with appearance fields solved on an approximately 8 µm grid. Each supplied class contributes weighted RGB sums S and masses m. A weak prior anchors the field to its global class mean μ. For class c, the sparse linear system is')
equation('(L_HE + diag(m_c + ρ)) F_c = S_c + ρ μ_c     (1)')
p('Here ρ = 0.015. L_HE is a four-neighbor graph Laplacian with nominal length scale ℓ = 16 µm. Edge conductance uses optical-density differences in the original H&E, with contrast τ = 0.1 and floor ε = 0.02:')
equation('w_pq = (ℓ / d_pq)² [ε + exp(−‖OD_p − OD_q‖² / (2τ²))]     (2)')
p('Distances d retain the physical axis spacing. Coarse class observations use a continuous tile-center weight ranging from 0.15 to 1 through a squared sine profile. Global class statistics use balanced capped samples and a pooled covariance with diagonal shrinkage and a numerical ridge. Sparse direct solves fit the three appearance coordinates. The fields are interpolated to the analysis grid; each class receives a negative Mahalanobis residual score against its fitted mean. Softmax produces relative competition scores and the highest score supplies the categorical ID. These scores are not calibrated semantic probabilities.')
p('The winning labels are projected directly from the analysis grid into the requested native output grid. H&E-derived exterior-background support is applied after refinement. No fine annotation, fixed margin, protected core, categorical median filter, watershed step or immutable seed is used in this AF path. This version therefore differs from the protected-boundary-wand algorithm in the supplied pilot manuscript. Other legacy Bora backends are not interchangeable with this method.')
hh('3.3 Bora VAF optional extension')
p('We name N08 Bora-VAF, Variance-Adaptive Appearance Fields. It whitens the evidence using the pooled covariance and replaces AF graph conduction with a squared-Laplacian bending operator. A second solve fits local log-variance from positive residual variance, using the implemented floor and shrinkage 0.2 + 0.8 max(residual, 0). Log-variance values are bounded to [−30, 30] for numerical stability, with clipping recorded. Gaussian residual and log-variance terms jointly determine competition. VAF is an optional comparison, not a claim that local Gaussian variance estimation is new.')
hh('3.4 Freeze anatomical regions before identification')
p('Within every anonymous group, eight-connected components define positive native region IDs. Every positive tissue pixel belongs to exactly one region; zero denotes excluded support. All 469 geometries, original image hashes, full-feature hashes and derived evidence hashes are frozen before any semantic target is attached. Separate files preserve region IDs and anonymous group IDs. Later classifiers may change a tissue name but cannot move, merge or redraw these boundaries.')
p('For each region, exact native-pixel overlap with token-grid cells supplies weighted full-feature means and standard deviations. Separate overlap weights quantify how much each physical tile contributes to that region. These features retain encoder context and are not independent cell measurements. A region may be heterogeneous or fragmented; the fixed partition does not guarantee anatomical correctness. The post hoc majority-reference accuracy ceiling measures the best possible pixel accuracy of one label per frozen region.')
hh('3.5 Five supervised region identification families')
p('We compare five distinct formulations. R01 aggregates the full ranked outputs of an existing native tile PLS-LDA head inside each region. Rank j among C classes receives Borda score (C − j + 1)/C, averaged using the exact tile-region overlap area. R02 fits PLS-LDA directly to each region’s 1,280-dimensional mean. R03 concatenates region mean and standard deviation. R04 regresses mapped annotated tissue fractions from the region mean and ranks the response predictions. R05 concatenates region mean with its difference from the shared-boundary-weighted neighboring mean; shared-edge lengths use native physical spacing. Isolated regions have zero neighbor contrast.')
p('Training targets are attached only after geometry freezing. A region is eligible for training or regional scoring when at least half its area has a mapped annotation. Its hard target is the dominant role within the annotated portion; ignored pixels are not converted to stroma. A stable reservoir caps training at 1,000 regions per dominant role, yielding 7,804 selected training segments and 25 represented roles. R02–R05 compare 16, 24 and 48 latent components with autoscaling. These component settings are hyperparameters within four families, not twelve different algorithms. The complete validation prediction grid is frozen before scoring. Region balanced accuracy selects each family configuration and the primary model.')
p('All PLS fits and predictions use the pinned fastPLS C++ CPU backend through Python bindings [10]. Native LDA rankings use predict with a top-five request; latent PLS scores are not interpreted as class scores. Composition regression ranks response predictions. The pinned simpls route is a bounded SIMPLS-family implementation and is not described as unqualified classical de Jong SIMPLS. Because the release lacks public fitted-object serialization, bundles preserve training arrays and controls and reconstruct the native head once. Startup refitting is separate from warm prediction latency. No per-image tissue vocabulary or expert mask is accepted by the region predictor.')
h('4 Experimental design')
hh('4.1 Cohort and annotation semantics')
p('The calibrated development corpus comprises 469 ROIs or WSI-derived windows from ten sources, not 469 complete-slide runs. It contains BCSS, TIGER, CAMELYON16, BEETLE, OCELOT, IGNITE, PESO, GCSS, PUMA and GleasonXAI [11–20]. Source releases, sample coordinates, calibration, native annotations and pairing hashes are retained in immutable inventories. Acquisition favored annotation-diverse fields and is not population-random sampling. WSSS4LUAD and RINGS fields with unresolved physical calibration do not enter this calibrated study. Gland-instance and cell-only datasets are outside this tissue-region endpoint.')
rows=[]
for dataset in sorted(registry['dataset_counts']):
    cases=[r for r in registry['cases'] if r['dataset']==dataset];counts=[sum(r['split']==s for r in cases) for s in ('train','validation','test')]
    qualifies=lambda r:len(set(r['ontology_raw_to_role'].values())-{0})>=2
    eligible=[sum(r['split']==s and qualifies(r) for r in cases) for s in ('validation','test')]
    rows.append([dataset, len(cases),*counts,*eligible])
rows.append(['Total',469,315,69,85,39,62])
table('Table 1. Internal cohort partitions and multiclass benchmark eligibility. All analysis units are ROIs or windows.',
      ['Source','All','Train','Val','Test','Val eligible','Test eligible'],rows,[1.65,.55,.6,.55,.55,.85,.85])
p('Twenty-seven stable broad tissue roles retain the original provider labels as metadata. Heterogeneous residual classes and automatic negative complements are excluded; non-cancer is not assumed to mean stroma. Gleason patterns collapse to tumor for this task and do not constitute several tissue groups. Single-role cases remain useful for training but are excluded from the multiclass benchmark. The mapping includes overlapping levels of tissue description and requires pathology review before clinical interpretation. PUMA’s unannotated real context is analyzed, while only the annotated center contributes reference scores.')
p('Known TCGA case identities and exact source-image duplicates are joined before an internal source-group partition, giving 315 training, 69 validation and 85 reserved images. Other source groups can be slides, cores or ROIs rather than verified patients, and unknown cross-source links remain possible. Provider data used for development are not untouched provider tests: PESO’s released test set is repartitioned into 24 internal training, 9 validation and 7 test fields. A source audit corrects the original registry’s overbroad provider-test exclusion statement without changing its immutable bytes. All images were previously inspected during Bora development, and the reserved cohort was already scored in earlier tile/superpixel experiments. The present follow-up preserves internal training-group separation but is exploratory repeated-test use.')
hh('4.2 Matched conditional refinement comparison')
p('The controlled comparison uses 39 eligible validation ROIs, identical reference-derived tile inputs, and the unchanged combined evidence RGB from algorithms 9 and 19. Fine annotations do not enter the refinement solver, but the starting classes and RGB are conditioned on coarse labels derived from those annotations. The endpoint measures conditional recovery from quantization and does not establish autonomous semantic performance. All methods share native-grid projection and H&E-derived background policy.')
p('Comparators are unrefined tiles, soft-observation H&E graph-label diffusion, and the authors’ Fast Bilateral Solver reference notebook [3]. The graph control is inspired by random walkers and quadratic smoothing; it is not an implementation of Grady’s original hard-seeded method. Bilateral parameters use an 8 µm spatial scale, luminance/chroma scales of 4, λ = 128 and relative conjugate-gradient tolerance 10⁻⁵. The initial 25-iteration results did not converge; a 500-iteration ceiling and six convergence-triggered reruns at 5,000 iterations yield convergence for every retained class solve. Compatibility adapters preserve the notebook definitions. This Python notebook is not the optimized implementation underlying the original paper’s speed claims.')
hh('4.3 Metrics timing and prespecified choices')
p('Native semantic-mask macro IoU averages over positive roles with nonzero union on annotated pixels. Spurious predicted roles are penalized; ignored annotation pixels do not enter quality measures. Per-image IoU is averaged with equal image weight. Annotated tissue removed by foreground estimation counts as error. Region top-one and top-five scores compare one ranking to each eligible region’s dominant expert role; balanced accuracy averages role recalls. These region metrics differ from pixel-area accuracy and cannot replace it.')
p('We use an Intel Core i7-13700 CPU with two numerical threads, low-priority workers and one numerical job at a time. Controlled refinement uses one warm-up and the median of three repeats, including guide construction, solving and native projection but excluding disk loading, encoder inference and RGB construction. A runtime penalty subtracts 0.05 times the mean excess above one second per ROI from mean IoU. Regional heads use five warm repeats; startup refit, evidence loading/context and geometry costs are recorded separately. Cached-stage sums are not measured end-to-end encoder throughput. Paired controlled differences use source-group means and 2,000 group bootstrap samples.')
p('R02 with 48 components is selected on validation and remains the primary follow-up classifier. All five validation-selected families predict all 85 reserved images before their new follow-up scoring; only 62 satisfy multiclass eligibility. A post hoc diagnostic ranks tissues solely by their training-region frequency. It uses no image features and does not alter model selection. Earlier tile and H&E superpixel studies are preserved as controls and are not represented as whole Bora segment classifiers.')
h('5 Results')
hh('5.1 Controlled refinement accuracy and speed')
table('Table 2. Matched cached refinement on 39 development ROIs. Bilateral results are convergence-qualified.',
      ['Method','Mean IoU','Mean seconds','ROIs above 1 s'],
      [[label,f"{af[name]['mean_iou']:.3f}",f"{af[name]['mean_seconds']:.3f}",af[name]['over_one_second']] for name,label in [('TILES','Tile input'),('GRAPH','Graph label diffusion'),('FBS-RGB','Bilateral reference'),('Bora-AF','Bora-AF'),('Bora-VAF','Bora-VAF')]],
      [2.65,1.1,1.25,1.35])
p('Bora-AF improves mean IoU from 0.586 to 0.663 and exceeds the converged bilateral reference’s 0.643 while requiring 0.168 versus 0.457 seconds per ROI. The source-group mean AF–bilateral difference is 0.0208 with a bootstrap 95% interval of 0.0093–0.0337. AF–tile improvement is 0.0795 with interval 0.0601–0.1018. The speed ratio is approximately 2.73 for these measured implementations; it is not a comparison with optimized bilateral code or a complete foundation-model pipeline. Graph diffusion is faster at 0.043 seconds but has lower IoU, 0.606.')
p('Bora-VAF reaches 0.662 IoU at 0.364 seconds. Its group-mean change relative to AF is −0.0004 with interval −0.0022 to 0.0015, providing no clear accuracy benefit in this cohort. Neither AF nor VAF exceeds one second on these controlled ROIs; five bilateral cases do. The runtime penalty therefore leaves AF and VAF unchanged and modestly reduces the bilateral score. These results support AF as the core method and VAF as an optional extension.')
figure('figure2-controlled-refinement.png','Figure 2. Conditional refinement distributions and mean cached CPU costs on the same 39 development ROIs. Diamonds show mean IoU, center lines show medians and boxes show interquartile ranges. Bilateral class solves all meet the specified convergence tolerance. This controlled task uses reference-derived coarse inputs and is distinct from autonomous region identification.')
hh('5.2 Identify whole frozen Bora segments')
p('Annotation-free geometry generation completes for all 469 images, producing 67,008 connected regions. Training uses 7,804 eligible selected regions. Validation scoring includes 4,016 segments across 39 images. R02-c48 has the highest validation region balanced accuracy, 47.2%, with 66.7% top-one accuracy and 98.9% top-five containment. Region feature dispersion and neighbor context give similar top-one results but do not change the selected primary model.')
rows=[]
labels={'R01':'Tile vote','R02':'Region mean','R03':'Mean and SD','R04':'Composition','R05':'Neighbor context'}
for key in ('R01','R02','R03','R04','R05'):
    a,b=val[key],test[key];rows.append([labels[key],f"{100*a['balanced_accuracy']:.1f}",f"{100*b['balanced_accuracy']:.1f}",f"{100*b['top1']:.1f}",f"{100*b['top5']:.1f}",f"{b['mean_native_pixel_iou']:.3f}",f"{1000*b['mean_head_seconds']:.2f}"])
b=priors['test'];rows.append(['Frequency control',f"{100*priors['validation']['balanced_accuracy']:.1f}",f"{100*b['balanced_accuracy']:.1f}",f"{100*b['top1']:.1f}",f"{100*b['top5']:.1f}",f"{b['mean_native_pixel_iou']:.3f}",'—'])
table('Table 3. Whole-region classification. Val and test balanced accuracy and top-one/top-five values are percentages. The primary model was fixed as region mean on validation.',
      ['Family','Val BA','Test BA','Test top 1','Test top 5','Test IoU','Head ms'],rows,[1.55,.70,.70,.80,.80,.80,.80])
p('The fixed primary classifier identifies 70.0% of 6,379 scored segments in 62 follow-up images and contains the correct dominant role in its top five for 96.7%; balanced accuracy is 49.0%. The training-frequency control obtains 35.3%, 92.3% and 6.25%, respectively. Thus much of the high top-five containment reflects common tissue roles, although image features improve it further. Rare-role recognition remains substantially weaker than overall accuracy suggests. One validation truth role is absent from hard-target training and is counted as error rather than removed.')
p('Pixel-mask IoU for the primary classifier is 0.232. Tile voting reaches 0.261 despite weaker region balanced accuracy, demonstrating that the best classifier depends on whether segments or pixel area carry equal importance. Region mean plus dispersion has the highest follow-up balanced accuracy, 51.0%, but is not promoted after inspecting test results. The post hoc perfect majority-label pixel-accuracy ceiling is 76.8% on the fixed follow-up geometry; mixed segments and support errors limit any single-label classifier on that partition.')
figure('figure3-segment-heads.png','Figure 3. Regional identity and dense-mask overlap favor different choices. Each family retains the configuration chosen on validation. The reused test is exploratory. The region-mean classifier remains the prespecified primary choice even though other families have higher follow-up scores on particular endpoints.')
p('The primary warm native head averages 1.91 milliseconds per eligible follow-up ROI. One standalone case requires a 3.74-second startup refit and then reproduces every saved region ranking and mask pixel without accepting an annotation. Average cached geometry costs are 1.17 seconds per eligible follow-up ROI, including grouping, RGB conversion, Bora, region pooling and recorded input/output work. Bora itself accounts for 0.176 seconds; pooling and component/output work account for approximately 0.670 seconds. These measurements prevent a millisecond classifier result from being mistaken for full-pipeline speed.')
figure('figure4-segment-example.png','Figure 4. Example tiger_62, selected by median primary-model IoU within the 16 eligible TIGER follow-up ROIs. Original H&E, expert roles, anonymous tile input, fixed RGB evidence, frozen Bora region geometry and region identities are shown. The geometry panel uses arbitrary region-ID colors; the identity panel and expert panel use the legend’s tissue colors. Black outlines retain separate regions even when their tissue labels agree. Every scale bar is 100 µm; masks are 42% opaque. Visible errors and mixed regions are retained. All 101 eligible validation and follow-up images are available in the companion viewers.')
h('6 Discussion')
p('The controlled results favor class appearance regularization over direct graph-label diffusion and the tested bilateral reference configuration. AF offers a useful accuracy–latency tradeoff for supplied coarse groups, while VAF adds computation without a demonstrated gain in this cohort. This evidence supports AF as the current core, but not a general claim of superior segmentation or speed across implementations, datasets or dense learned models.')
p('The semantic extension now answers a regional question. A tile classifier can contribute evidence, but its output is aggregated inside a Bora segment rather than displayed as the final unit. Direct regional pooling improves region identity relative to tile voting, while the reverse ordering for pixel IoU reveals unresolved area and geometry effects. High top-five containment is not equivalent to a reliable mask, calibrated confidence or successful rare-tissue recognition. Mixed-region ceilings and frequency controls make these limitations visible.')
p('The main limitations are repeated cohort inspection, boundary-enriched acquisition, imperfect source-group reconciliation, sparse annotations, ontology overlap and the fixed four-group partition. Foreground estimation can mistake pale border tissue for exterior glass; its annotated errors remain penalized. Physical calibration controls encoder context but does not recover optical resolution, particularly in interpolated exports. Native output dimensions do not imply independent native-scale feature information. Neither the ROI timings nor cached feature experiments establish complete-WSI throughput.')
p('The strongest next steps are an independent external cohort, a pathology-reviewed hierarchical ontology, adaptive partitioning with explicit mixed-region handling, and matched training/evaluation of a compact dense student. Learned feature upsampling or prompt-derived geometry may be useful components [5–9], but their accuracy and cost must be measured under the same input and output contracts. Regional confidence calibration and abstention would also be needed before treating ranked choices as decision support. These extensions are proposals, not results of this study.')
h('7 Conclusion')
p('Bora-AF provides efficient conditional refinement by fitting H&E-guided class appearance fields, with an optional variance-adaptive comparison. A separate supervised stage can identify complete frozen Bora regions using native CPU PLS models and preserve a top-five tissue shortlist for each segment. Controlled refinement is supported by matched development comparisons; autonomous semantic performance remains exploratory and limited by partition quality and tissue-role generalization.')
h('Data code and reproducibility')
p('Public source records and acquisition identities are retained in the benchmark inventories. Bora code is maintained at github.com/tkcaccia/Bora; the pinned fastPLS binding and canonical C++ core are recorded with the run manifests [10]. The experiment modules are bora_supervised and bora_segment. All geometry, target preparation, candidate predictions, model controls and source revisions have immutable hashes. Native region IDs, semantic masks and regional rankings are separate outputs. The registry can append images and role IDs while preserving prior identities, mappings and internal splits. A new model and dataset version is required for retraining; an inspected test cohort is not silently relabeled as a fresh test.')
p('Source images and dataset licenses remain with their providers. The local and Chiamaka research corpus contains the full output evidence; the manuscript does not claim that all large derived arrays or source images are already available as a public archival release. The original pilot manuscript is preserved unchanged. Authors must confirm affiliations, corresponding author, contributions, funding, conflicts and the institutional determination for public-data reuse before submission. These declarations were not supplied and have not been invented.')
h('References')
refs=[
('Zimmermann E, Vorontsov E, Viret J, et al. Virchow2 Scaling Self Supervised Mixed Magnification Models in Pathology. 2024. arXiv 2408.00738.','https://arxiv.org/abs/2408.00738'),
('Grady L. Random walks for image segmentation. IEEE Transactions on Pattern Analysis and Machine Intelligence. 2006;28:1768–1783.','https://pubmed.ncbi.nlm.nih.gov/17063682/'),
('Barron JT, Poole B. The Fast Bilateral Solver. European Conference on Computer Vision. 2016. arXiv 1511.03296.','https://arxiv.org/abs/1511.03296'),
('Wang L, He L, Mishra A, Li C. Active contours driven by local Gaussian distribution fitting energy. Signal Processing. 2009;89:2435–2447.','https://doi.org/10.1016/j.sigpro.2009.03.014'),
('Fu S, Hamilton M, et al. FeatUp A Model Agnostic Framework for Features at Any Resolution. ICLR. 2024.','https://arxiv.org/abs/2403.10516'),
('Seo J, Hamilton M, Kim S. Upsample Anything A Simple and Hard to Beat Baseline for Feature Upsampling. arXiv 2511.16301.','https://arxiv.org/abs/2511.16301'),
('AnyUp Universal Feature Upsampling. arXiv 2510.12764. Research preprint.','https://arxiv.org/abs/2510.12764'),
('Liu et al. SegmentAnyTissue. Medical Image Analysis. 2025;102:103550.','https://doi.org/10.1016/j.media.2025.103550'),
('Lv J, Zhu Y, Okada S, et al. Computationally Efficient Pathology Segmentation using Knowledge Distillation from Foundation Models. 2026. arXiv 2609.03947v1. Submitted to Medical Image Analysis; preprint.','https://arxiv.org/abs/2609.03947'),
('Cacciatore S. fastPLS and fastPLS py. C++ numerical core and Python bindings. Version 0.3.0; binding commit fc0c011969b754ee196aacc220a217783afeb05d; core 82bbc48a0d69e4bd0d7c261fdcc8e636694133b1.','https://github.com/tkcaccia/fastPLS-py'),
('Amgad M, Elfandy H, Hussein H, et al. Structured crowdsourcing enables convolutional segmentation of histology images. Bioinformatics. 2019;35:3461–3467. BCSS source and pinned mirror recorded in the inventory.','https://doi.org/10.1093/bioinformatics/btz083'),
('TIGER Grand Challenge. TIGER WSIROIS data and tissue annotations. Zenodo record 6014422.','https://zenodo.org/records/6014422'),
('CAMELYON16. Original whole-slide source and tumor annotations. Registry of Open Data on AWS.','https://registry.opendata.aws/camelyon/'),
('Lems CM, et al. BEETLE multicentric breast tissue segmentation dataset. Zenodo record 16812932.','https://zenodo.org/records/16812932'),
('OCELOT. Original challenge data and author repository.','https://github.com/lunit-io/ocelot23algo'),
('IGNITE version 2. H&E tissue segmentation data. Zenodo record 17735903.','https://zenodo.org/records/17735903'),
('PESO test release version 1.1. Zenodo record 5137717. This release is reused as development data here.','https://zenodo.org/records/5137717'),
('GCSS. Gastric cancer tissue segmentation data. Zenodo record 8176354.','https://zenodo.org/records/8176354'),
('PUMA. Training tissue annotations and real image context. Zenodo record 15050523.','https://zenodo.org/records/15050523'),
('GleasonXAI. TissueArray images and region annotations. Figshare version 1.','https://doi.org/10.6084/m9.figshare.27301845.v1')]
for number,(text,url) in enumerate(refs,1):reference(number,text,url)
doc.core_properties.title='Bora appearance fields for tissue segmentation and segment identification'
doc.core_properties.subject='Working article draft with conditional refinement and whole region identification'
doc.core_properties.author='Bora research team'
output=ROOT/'Bora_Appearance_Fields_MIA_draft_2026-10-18.docx';doc.save(output)
provenance=dict(output=str(output),original_preserved=str(ROOT.parent/'source-chiamaka-2026-10-18'),
    supervised_unit='Connected native Bora tissue segment',primary_model='R02-c48',
    source_reports={str(path):hashlib.sha256(path.read_bytes()).hexdigest() for path in [SEG/'REPORT.json',SEG/'classifiers-validation-v1/RESULTS.json',SEG/'classifiers-followup-test-v1/RESULTS.json',OLD/'CONTROLLED_FINAL.json',OLD/'REGISTRY.json']},
    authorship_source='Original manuscript; verification notes retained',unsubmitted_working_draft=True)
(ROOT/'MANUSCRIPT_PROVENANCE.json').write_text(json.dumps(provenance,indent=2)+'\n')
print(output)
