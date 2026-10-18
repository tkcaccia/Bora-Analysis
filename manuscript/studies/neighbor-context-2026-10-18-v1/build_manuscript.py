"""Versioned article and editorial assessment from frozen Bora evidence."""
from pathlib import Path
from collections import Counter
from copy import deepcopy
import json,hashlib,re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
BENCH=ROOT.parents[1];REPORTS=BENCH/'results'
GROUP=REPORTS/'bora-anonymous-groups-2026-10-18-v1'
EVO=REPORTS/'bora-classifier-evolution-2026-10-18-v2'
SOURCE=ROOT.parent/'source-chiamaka-2026-10-18/Bora_MIA_pilot_manuscript.docx'
load=lambda p:json.loads(p.read_text())
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
view=load(GROUP/'viewer-no-puma-area20-v7/data.json')
registry=load(ROOT/'evidence/REGISTRY.json')
cohort=load(EVO/'COHORT_REVISION.json')
controlled=load(REPORTS/'bora-supervised-2026-10-18-v1/CONTROLLED_FINAL.json')
audit=load(GROUP/'geojson-no-puma-area20-v7/VALIDATION.json')
ref=load(ROOT/'evidence/REFERENCE_SUMMARY.json')
assert len(view['cases'])==51 and cohort['images']==116
assert all(c['dataset']!='puma' for c in view['cases']) and cohort['puma_evaluation_images']==0
assert all(audit[k]==0 for k in ['positive_polygon_parts_below20','polygon_holes_below20','positive_raster_regions_below20','enclosed_raster_holes_below20','interclass_overlap_area_pixels'])


def figures():
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False})
    # Matched results are complete, unlike the ongoing expanded classifier study.
    fig,ax=plt.subplots(1,2,figsize=(9,3.25),layout='constrained')
    names=['TILES','GRAPH','FBS-RGB','Bora-VAF','Bora-AF'];labels=['Input tiles','Label diffusion','Bilateral reference','Bora VAF','Bora AF']
    lookup={r['name']:r for r in controlled['ranking']};colors=['#919ca2','#919ca2','#919ca2','#7ca3b3','#176581']
    ax[0].barh(labels,[lookup[n]['mean_iou'] for n in names],color=colors)
    ax[0].set(xlim=(0,.75),xlabel='Mean native mask macro IoU',title='A  Matched conditional refinement  n = 39')
    for i,n in enumerate(names):ax[0].text(lookup[n]['mean_iou']+.008,i,f"{lookup[n]['mean_iou']:.3f}",va='center',fontsize=8)
    ax[1].barh(labels,[lookup[n]['mean_seconds'] for n in names],color=colors)
    ax[1].set(xlim=(0,.55),xlabel='Mean seconds per ROI',title='B  Measured refinement stage only')
    for i,n in enumerate(names):ax[1].text(lookup[n]['mean_seconds']+.008,i,f"{lookup[n]['mean_seconds']:.3f}",va='center',fontsize=8)
    fig.savefig(ROOT/'figures/controlled.png',dpi=240);plt.close(fig)
    # A median example and the worst available example make selection explicit.
    ordered=sorted(view['cases'],key=lambda c:next(m for m in c['models'] if m['family']=='R05')['metrics']['pixel_score']['macro_iou'])
    examples=[ordered[len(ordered)//2],ordered[0]]
    from matplotlib.patches import Patch
    fig,axs=plt.subplots(4,3,figsize=(6.65,7.1),layout='constrained')
    titles=['Original H&E','Expert tissue mask','Anonymous input tiles','Unchanged RGB 9 + 19','Neighbor context output']
    for row,c in enumerate(examples):
        for col,file in enumerate(['original','expert','tiles','converted','R05-c48']):
            axis=axs[2*row+col//3,col%3]
            im=Image.open(GROUP/'viewer-no-puma-area20-v7'/c['case']/(file+'.jpg'));axis.imshow(im)
            axis.set_xticks([]);axis.set_yticks([])
            for spine in axis.spines.values():spine.set_visible(False)
            axis.set_title(titles[col],fontsize=8)
            # Preview coordinates stretch native x/y identically to the viewer.
            bar=100/c['native_mpp_xy'][0]*im.width/c['native_shape_yx'][1]
            x=.06*im.width;y=.91*im.height
            axis.plot([x,x+bar],[y,y],color='white',lw=5);axis.plot([x,x+bar],[y,y],color='black',lw=2)
            axis.text(x,y-.035*im.height,'100 µm',color='black',fontsize=7,bbox=dict(facecolor='white',alpha=.8,edgecolor='none',pad=1))
        m=next(m for m in c['models'] if m['family']=='R05')
        axs[2*row,0].set_ylabel(c['case']+'\nFinal IoU '+f"{m['metrics']['pixel_score']['macro_iou']:.3f}",fontsize=8)
        roles={g['role'] for g in c['expert_legend'] if g['role']>0}|set(m['top1'])
        handles=[Patch(facecolor=np.asarray(view['colours'][g])/255,label=view['roles'][g-1].replace('_',' ')) for g in sorted(roles)]
        axs[2*row+1,2].axis('off');axs[2*row+1,2].legend(handles=handles,loc='center left',frameon=False,fontsize=8,handlelength=1)
        axs[2*row+1,2].set_title('Tissue colors',fontsize=8)
    fig.savefig(ROOT/'figures/examples.png',dpi=240);plt.close(fig)
    return [c['case'] for c in examples]


def new_document(title,subtitle):
    # Retain the source article's page geometry and style definitions; replace
    # obsolete body content rather than editing the author's original file.
    doc=Document(SOURCE)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype==RT.IMAGE:doc.part.drop_rel(rid)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        fonts=s._element.get_or_add_rPr().get_or_add_rFonts();fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.09;normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',19),('Heading 1',13),('Heading 2',11)]:
        s=doc.styles[name];s.font.size=Pt(size);s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(10 if name!='Title' else 0);s.paragraph_format.space_after=Pt(5)
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False
    for section in doc.sections:
        for header in (section.header,section.first_page_header,section.even_page_header):
            for p in header.paragraphs:p.text=''
        section.header.paragraphs[0].text='Bora appearance fields and Neighbor context  |  Working draft 18 October 2026'
        for r in section.header.paragraphs[0].runs:r.font.size=Pt(8)
        footer=section.footer.paragraphs[0];footer.text='';footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_paragraph(title,'Title');p=doc.add_paragraph(subtitle);p.paragraph_format.space_after=Pt(9)
    for r in p.runs:r.italic=True;r.font.size=Pt(9)
    return doc


def para(doc,text,style=None):return doc.add_paragraph(text,style)
def heading(doc,text,level=1):return doc.add_heading(text,level=level)
def equation(doc,text):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
    def run(parent,value,literal=False):
        r=OxmlElement('m:r')
        if literal:
            props=OxmlElement('m:rPr');lit=OxmlElement('m:lit');lit.set(qn('m:val'),'1');props.append(lit);r.append(props)
        t=OxmlElement('m:t');t.text=value;r.append(t);parent.append(r)
    cursor=0
    for match in re.finditer(r'([A-Za-zμνΣ]+)_([A-Za-z]+)',text):
        run(m,text[cursor:match.start()]);sub=OxmlElement('m:sSub');base=OxmlElement('m:e');index=OxmlElement('m:sub')
        run(base,match[1])
        for letter in match[2]:run(index,letter,True)
        sub.append(base);sub.append(index);m.append(sub);cursor=match.end()
    run(m,text[cursor:]);p._p.append(m)
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
def table(doc,caption,headers,rows,widths):
    p=para(doc,caption,'Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for cell,value,w in zip(t.rows[0].cells,headers,widths):cell.text=str(value);cell.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell,value,w in zip(t.add_row().cells,values,widths):cell.text=str(value);cell.width=Inches(w)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;props=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            props.append(b);m=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
            props.append(m)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EBEF');props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9);r.bold=i==0
    para(doc,'').paragraph_format.space_after=Pt(2)
def figure(doc,path,caption):
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Inches(6.65))
    para(doc,caption,'Caption')
def reference(doc,n,text,url):
    p=para(doc,f'[{n}] {text} ');p.paragraph_format.space_after=Pt(2)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=url;r.append(t);link.append(r);p._p.append(link)
    for r in p.runs:r.font.size=Pt(9)


def article(examples):
    doc=new_document('Bora appearance fields for refinement and identification of coarse tissue groups in digital pathology',
        'Article draft integrating Neighbor context and compact GeoJSON  |  Completed development evidence and an ongoing expanded study')
    source=Document(SOURCE)
    for p in source.paragraphs:
        if p.text.startswith('Authors:') or p.text.startswith('¹ Bioinformatics'):para(doc,p.text)
    heading(doc,'Abstract')
    para(doc,'Tile-level pathology maps provide coarse tissue assignments whose square boundaries obscure anatomical interfaces. We present Bora-AF, H&E-guided Appearance Fields, for conditional refinement of a supplied multiclass map, followed by supervised identification of its anonymous tissue groups. Bora-AF fits spatial class appearances on a small physically calibrated graph and applies categorical competition over the tissue field. Neighbor context pools full Virchow2 features within each refined group, appends its contrast to adjacent groups, and identifies tissue through a native fastPLS CPU classifier. A shared-boundary vector stage replaces staircase contours with calibrated straight chords and removes polygon fragments and enclosed holes below 20 native pixels of area. In a completed 39-ROI development comparison, Bora-AF achieved mean native-mask macro IoU of 0.663 in 0.168 seconds per ROI; a convergence-qualified bilateral reference achieved 0.643 in 0.457 seconds. These timings exclude the encoder and evidence construction. On 51 previously inspected PUMA-free test ROIs or windows, Neighbor context identified 91.3% of 173 anonymous groups, with 70.4% balanced accuracy and mean final GeoJSON IoU of 0.606. The source-group bootstrap interval for its IoU difference from pooled-mean classification crossed zero. All 255 reference exports passed continuous-area and native-raster cleanup checks without interclass overlap. An expanded 116-image development evaluation is in progress and contributes no final results here. The findings support efficient conditional refinement and group identification; expert-derived coarse inputs and adaptive test reuse preclude claims of autonomous or independently validated semantic segmentation.')
    para(doc,'Keywords: computational pathology; tissue segmentation; appearance fields; partial least squares; spatial context; GeoJSON','Caption')
    heading(doc,'1 Introduction')
    para(doc,'Quantifying tumor, stroma, immune infiltrates and other tissue compartments requires both a geometric partition and an interpretable tissue identity. Pathology foundation models such as Virchow2 supply informative representations [1], but projecting one categorical assignment per tile gives grid-shaped contours. A tile can also contain several compartments. Naming a tile and locating the borders of the complete tissue group are therefore related but distinct operations.')
    para(doc,'Bora addresses the conditional problem in which an aligned image and a coarse categorical map already exist. We first use image-guided class appearance estimates to refine the supplied groups, then identify each resulting anonymous group from its pooled foundation-model features and physical neighborhood. Separating these operations permits the same geometry to be inspected independently of the classifier. A later vector stage produces compact native-coordinate polygons suitable for overlays and morphometric analysis.')
    para(doc,'Bora-AF refines the complete tissue field using soft coarse observations. Its spatial regularizer acts on class appearance fields; it does not restrict label changes to a prescribed margin or make a seed interior immutable. In the controlled experiment, the solver receives anonymized expert-derived tile assignments. This experiment measures recovery from coarse expert-derived structure. It does not establish how well independently predicted clusters would segment previously unseen clinical material.')
    para(doc,'The contributions investigated are the combination of H&E-conducted appearance fields for full-field multiclass competition, an explicit whole-group identification stage using native PLS and boundary-weighted context, and a shared-border vectorization contract that controls geometric complexity in physical units. Graph regularization, local Gaussian fitting, PLS and straight-line simplification are established techniques. The contribution is the specified pathology workflow and its measurable behavior; priority or superiority over all related methods is not claimed.')
    heading(doc,'2 Related work')
    para(doc,'Image-guided segmentation has a long history. Random walkers infer labels from seeded graph structure [2]. The Fast Bilateral Solver provides confidence-weighted edge-aware optimization [3]. Local Gaussian distribution fitting estimates spatially varying means and variances during contour evolution [4]. These precedents limit novelty claims for smoothing, graph Laplacians or local appearance statistics alone. Bora-AF instead holds coarse observations fixed while fitting class appearance fields, then compares every group against those fields. This is a particular conditional estimator, rather than a new general principle of local statistical segmentation.')
    para(doc,'FeatUp restores spatial detail to learned representations using image guidance and multiview consistency [5]. Bora works downstream with a categorical coarse map and a fixed three-channel evidence image. The two approaches may be complementary, but no matched FeatUp comparison has been completed. A high-resolution output array also does not establish high-resolution semantic information: the encoder sampling, field grid, output projection and optical acquisition each impose separate limits.')
    para(doc,'Lv and colleagues train Virchow2-based segmentation teachers and distill response and feature knowledge into compact dense students [6]. Their objective overlaps with ours in efficient tissue segmentation. Their student predicts dense tissue labels from the image after supervised training, whereas the present Bora route requires an explicit coarse map and Virchow2 features or their cache. Our CPU head identifies groups after geometry has been formed. It is not a distilled dense decoder. Faster warm classification or cached refinement cannot establish faster deployment than a compact student, because encoder and evidence construction costs remain to be measured under matched conditions.')
    para(doc,'Native fastPLS supplies the numerical classification engine [7]. Adding a boundary-weighted neighborhood contrast is a deliberately compact contextual representation, not a claim to invent contextual classification. Its contribution must be judged by comparisons with the same PLS head without context and with stronger spatial and dense-prediction baselines. The currently completed comparison is exploratory.')
    heading(doc,'3 Methods')
    heading(doc,'3 1 Inputs physical calibration and controlled tiling',2)
    para(doc,'Inputs are an H&E image, an aligned categorical tile mask, a tissue support estimate, a three-channel feature evidence image and native and working pixel spacing for both axes. Image fields are designated as ROIs or WSI-derived windows; their source WSI status is not interpreted as evidence of complete-slide processing. Native spacing is retained through integer resizing, so the effective working spacing can differ slightly from its nominal value of 0.5 µm per pixel.')
    para(doc,'At the nominal working grid, a 224-pixel Virchow2 window covers 112 µm and its 14-pixel spatial-token interval corresponds to 7 µm. These are sampling and contextual field sizes, not optical resolving power. The encoder is pinned to revision 3158645804b69e3f3bc4439d4116edddf0840a72. Full 1,280-channel spatial-token arrays are retained. Unknown calibration, insufficient physical field size, insufficient usable tiles and fewer than two mapped tissue groups exclude an image from the multiclass evaluation. Acquisition requires at least four complete encoder windows per axis and 16 usable coarse observations; overlap or upsampling is not used to manufacture eligibility.')
    para(doc,'For this controlled benchmark, majority expert labels define the previous physical tile assignments, using approximately 45 µm fields, or 90 pixels at the nominal working scale. A deterministic permutation replaces class numbers with anonymous IDs without changing assignments. A segregated numeric key stores the source identity for training and evaluation. The solver receives the anonymous map, and the classifier receives pooled image features and geometric context. Source names, the answer key, semantic palette and per-image reference vocabulary are excluded from the classifier inputs. An independent solve with the original numeric IDs verifies that anonymization changes no refined pixel after restoring the labels.')
    heading(doc,'3 2 Fixed Virchow2 evidence conversion',2)
    para(doc,'The evidence conversion combines scale-persistent evidence and predictable latent structure, designated algorithms 9 and 19 in the reproducible implementation, with equal weights. The resulting RGB is frozen before the classifier comparison. Both methods derive within-image class evidence using coarse tile groups, so the evidence is reference-conditioned in this benchmark even though neither method reads the fine mask during prediction. Their output is a synthetic representation of model evidence, not reconstructed H&E, a microscopy channel or a molecular measurement.')
    para(doc,'The conversion first selects 64 feature coordinates using coarse group discrimination. Algorithm 9 emphasizes coordinates whose tile summaries remain consistent across nominal 0.5, 1 and 2 µm samplings; 32 balanced discriminative coordinates contribute to its cross-scale consensus. Algorithm 19 predicts selected latent coordinates from neighboring summaries at offsets of one, two and four tokens using ridge regression. Four spatial folds use a four-token exclusion buffer, although encoder receptive fields still overlap. Aligned group-evidence probabilities are averaged before palette mapping and eight-bit RGB quantization. The semantic classifier subsequently uses all 1,280 original feature coordinates rather than only the conversion-selected subset.')
    heading(doc,'3 3 H and E guided appearance fields',2)
    para(doc,'Let R denote the evidence RGB, C the supplied anonymous group map and H the original H&E. Evidence is analyzed on an approximately 1 µm grid. Each group c contributes a spatial observation mass m_c and weighted RGB sum S_c, reduced to an approximately 8 µm field grid. A continuous tile-center confidence ranges from 0.15 to 1 with a squared-sine profile. It weights observations without making a protected core immutable. A weak global appearance prior has strength ρ = 0.015.')
    para(doc,'On a four-neighbor field graph, edge conductance depends on optical-density contrast in H and the physical distance d_pq. With length scale ℓ = 16 µm, contrast τ = 0.1 and conductance floor ε = 0.02, the implemented weights are')
    equation(doc,'w_pq = (ℓ / d_pq)² [ε + exp(−‖OD(H_p) − OD(H_q)‖² / (2τ²))]     (1)')
    para(doc,'The associated Laplacian L_H regularizes the class appearance F_c. For each RGB coordinate, the discrete objective has a data term m_c F_c² − 2 S_c F_c and prior ρ(F_c − μ_c)², together with the graph quadratic. Its stationary equation is')
    equation(doc,'[L_H + diag(m_c + ρ)] F_c = S_c + ρ μ_c     (2)')
    para(doc,'The fixed-observation field problem is strictly convex because ρ is positive and the graph Laplacian is positive semidefinite. Sparse direct solves fit the three appearance coordinates. This property does not make the subsequent categorical segmentation globally convex or guarantee anatomically correct borders. The positive conductance floor permits information to pass weak stain edges, while strong H&E contrast reduces borrowing across an interface.')
    para(doc,'Balanced capped coarse samples estimate global class means and a pooled covariance. The covariance uses 0.8 times its empirical value, 0.2 times its diagonal and a numerical ridge of 0.0025 in standardized evidence units. Fitted fields are interpolated to the analysis grid. Every group competes everywhere in the supported tissue through its negative Mahalanobis residual against the local field. Softmax supplies relative competition scores; the largest score gives the categorical group. These scores are not calibrated tissue probabilities. Labels are projected directly to the requested native grid, then H&E-derived exterior background is removed. Unannotated real tissue remains available for feature extraction and refinement; it is excluded only from reference quality calculations.')
    heading(doc,'3 4 Optional variance adaptive extension',2)
    para(doc,'Bora-VAF is retained as a separate comparison. It whitens the evidence, uses a squared-Laplacian bending operator, and fits a local log-variance field in a second solve. The variance target applies the implemented floor and shrinkage 0.2 + 0.8 max(residual, 0); bounded log-variance terms enter Gaussian competition. This increases computation and is not the primary AF method. Local variance modeling itself is established [4].')
    heading(doc,'3 5 Whole anonymous tissue groups',2)
    para(doc,'The classification unit is one surviving anonymous Bora group, including all its disconnected pieces. Connected components are not separately classified. Before semantic training, the native group plane, token and tile overlap matrices, evidence image, pooled features and all source hashes are frozen. Later classifiers can change tissue identity but cannot change these AF boundaries. Different anonymous groups may receive the same tissue label; their provenance IDs remain recorded even if the final per-class GeoJSON dissolves their adjoining geometry.')
    para(doc,'For group g and spatial token t, exact native-pixel overlap supplies an area weight a_gt. The group mean retains every encoder coordinate:')
    equation(doc,'μ_g = (Σ_t a_gt z_t) / (Σ_t a_gt)     (3)')
    para(doc,'The supervised hard target is the hidden source tissue identity that generated the original coarse group. It is not a post-refinement fragment-majority label. Training and scoring require a mapped source identity and at least one mapped reference pixel. Unannotated pixels do not become stroma or a negative tissue class. The coarse partition and RGB remain expert-conditioned, so freezing the feature pipeline prevents subsequent target access but does not turn this task into autonomous segmentation.')
    heading(doc,'3 6 Neighbor context classification with native fastPLS',2)
    para(doc,'Neighbor context represents a group by its own mean and its contrast to neighboring groups. Let b_gh be the length of its shared edge with group h, measured using the appropriate native pixel spacing for each axis. Exterior background supplies no neighbor. The neighboring mean and classification vector are')
    equation(doc,'ν_g = (Σ_h b_gh μ_h) / (Σ_h b_gh) ;     x_g = [μ_g , μ_g − ν_g]     (4)')
    para(doc,'An isolated group uses ν_g = μ_g, giving zero contrast. The resulting vector has 2,560 features. A native fastPLS CPU PLS-LDA model with autoscaling supplies a tissue ranking. The present reference configuration uses 48 components, selected within the Neighbor context family on validation, with seed 314159. Its original training set contains 838 whole groups from training images and 25 represented tissue roles. A per-role cap of 1,000 is inactive at this group count. The pinned simpls implementation is reported by library revision rather than being presented as a new PLS algorithm.')
    para(doc,'The described family is provisional while the expanded study evolves. The historic pooled-mean family was the validation-selected primary reference; choosing Neighbor context for the present manuscript after inspecting development results is a post hoc methodological choice. The reference model is not retrained on test labels. Because the installed binding lacks a public fitted-object serializer, reproducible bundles retain training arrays and controls and refit once at startup. Startup, feature pooling and neighborhood construction are separate from warm head prediction latency.')
    heading(doc,'3 7 Compact shared GeoJSON geometry and final cleanup',2)
    para(doc,'The published Bora coverage converter remains the primary vectorizer, using the accurate profile with a 500-native-pixel component filter and two-pixel shared-border simplification. Those converter settings are distinct from the final 20-pixel cleanup. Per-class polygons are repaired and dissolved while sharing interface geometry. A subsequent shared-boundary graph is calibrated along both axes in micrometres. Maximal arcs stop at junctions and are fitted once using Ramer–Douglas–Peucker straight chords. Polygonizing the fitted network produces mutually consistent faces, which receive the class with greatest overlap with the preceding predicted partition. No expert reference enters this assignment or fitting.')
    para(doc,'The requested tolerance is at most 3 µm. Lower tolerances are tried until every predicted class is preserved, each class has at least 0.97 shape IoU with the preceding predicted geometry, and the tissue-change fraction is at most 0.015. The latter counts half the summed per-class symmetric differences plus half the exterior-support symmetric difference, divided by original predicted tissue area. These are fidelity guards against a predicted geometry, not annotation-accuracy guarantees. A coverage failing the fitting guards retains its preceding geometry.')
    para(doc,'Final cleanup enforces both continuous polygon-part area and native raster counts. Positive polygon parts strictly smaller than 20 native pixels of area are removed and assigned to adjacent tissue with the longest shared border; an isolated speck on glass is removed to background. Enclosed empty holes below the same threshold are filled. A raster check repairs remaining connected native fragments and enclosed holes, using only affected local pixel-cell patches. Geometric and raster checks repeat until both rules hold. Exactly-threshold areas are retained by this final cleanup; the earlier 500-pixel vectorizer filter remains separately documented. Empty exterior glass is not dilated into tissue. Final quality is calculated after rasterizing these actual polygons.')
    heading(doc,'4 Experimental design')
    heading(doc,'4 1 Dataset accounting and scope',2)
    para(doc,'The expanded frozen corpus has 533 ROIs or WSI-derived windows, partitioned into 349 training, 72 validation and 112 test images. It extends the 469-image registry with 40 GCSS and 24 IGNITE fields from previously unused source groups. The active evaluation excludes PUMA from both reserved splits without moving excluded images into training. PUMA remains part of the training corpus. This evaluation-scope change followed earlier development inspection and is post hoc. Single-role sources can supply training examples but do not qualify for multiclass evaluation. The current ontology contains 28 stable, append-only roles; the trained heads represent 25 of them. An unseen reference role is counted as error rather than silently omitted.')
    ds_rows=[]
    ids=set(cohort['cases'])
    for dataset in sorted(registry['dataset_counts']):
        rs=[r for r in registry['cases'] if r['dataset']==dataset]
        counts=[sum(r['split']==s for r in rs) for s in ('train','validation','test')]
        ev=[sum(r['split']==s and r['case'] in ids for r in rs) for s in ('validation','test')]
        ds_rows.append([dataset,len(rs),*counts,*ev])
    ds_rows.append(['Total',533,349,72,112,38,78])
    table(doc,'Table 1. Frozen image inventory and active evaluation eligibility. Counts refer to ROIs or windows, not independent patients or complete WSI runs.',
        ['Source','All','Train','Val','Test','Val eligible','Test eligible'],ds_rows,[1.6,.58,.6,.58,.58,.85,.85])
    para(doc,'The active cohort has 116 images: BCSS42, TIGER26, GCSS24 and IGNITE24 [8–11]. It contains 78 test and 38 validation images; test coverage increases from 51 non-PUMA fields in the inspected reference viewer to 78, adding 18 IGNITE and nine GCSS test fields. The original five-method reference results in this draft use the existing 51-image subset only: BCSS27, TIGER16, GCSS5 and IGNITE3. Their 173 evaluated groups arise from 45 recorded source groups. Source groups can be slides or specimens rather than verified patients.')
    para(doc,'BCSS and TIGER are not automatically independent: the official TIGER release includes adapted BCSS-derived TCGA annotations [9]. Known TCGA donors and exact source-image duplicates are linked before partitioning. Unknown cross-source relationships remain a limitation. Heterogeneous residual labels and automatic non-tumor complements are not relabeled as stroma. Gleason grades collapse to their shared tumor role for this tissue endpoint. Acquired WSSS4LUAD and RINGS fields without verified physical calibration do not enter this calibrated evaluation. Gland-instance and cell-only endpoints are not pooled with tissue-compartment segmentation.')
    para(doc,'The full training inventory also draws on released CAMELYON16, BEETLE, OCELOT, PESO, PUMA and GleasonXAI images and region annotations [12–17]. Their provider partitions and annotation semantics are retained in the source audit. In particular, the acquired PESO provider test release has already been reused as internally repartitioned development data; it is not an untouched provider test. No claim of universal provider-test isolation is made. A source-group separation within this registry does not restore the independence lost through earlier development use.')
    heading(doc,'4 2 Completed refinement comparison',2)
    para(doc,'The matched development comparison contains 39 eligible validation ROIs with identical expert-derived tile inputs, frozen RGB9+19 evidence and background policy. Methods are input tiles, soft-observation H&E graph-label diffusion, the authors’ bilateral-solver reference notebook, Bora-AF and Bora-VAF. The graph control is not an implementation of the original hard-seeded random walker. Bilateral parameters use an 8 µm spatial scale, luminance and chroma scales of 4, λ = 128 and relative conjugate-gradient tolerance 10⁻⁵. A convergence audit retains already-converged outputs and reruns six initially nonconverged cases at a 5,000-iteration ceiling; every retained class solve converges. This notebook comparison does not test the optimized implementation underlying the bilateral paper’s throughput claims.')
    heading(doc,'4 3 Classifier development and ongoing evolution',2)
    para(doc,'The completed reference compares tile-rank voting, pooled mean, mean plus dispersion, tissue-composition regression and Neighbor context on the same frozen group geometry. Family component counts were chosen on validation. PUMA-free aggregate scores are recalculated from the retained cases after the latest vector cleanup. A paired source-group bootstrap of the Neighbor context versus pooled-mean final IoU difference averages cases within each of 45 source groups and resamples groups 2,000 times, using seed 314159. Its interval describes the inspected cohort and does not adjust for adaptive method selection.')
    para(doc,'The ongoing expansion trains on 967 whole groups. Each of ten generations evaluates 20 fastPLS-based mechanisms, retains five and generates three structural children per retained parent. Initial mechanisms include group discriminants, covariance and quantile descriptions, tile evidence pooling, hierarchical and pairwise classification, coded responses, kernel PLS, OPLS, residual models, source bagging, local experts, spatial pooling, composition transport, latent densities and verification metrics. Supporting PCA and clustering do not replace the required native PLS core. Candidate predictions for all 116 images freeze before fine-mask scoring within a generation. Structural children use training-source cross-fitting where teacher evidence is needed.')
    para(doc,'The selection score is 0.50 times image-mean final GeoJSON IoU plus 0.30 times group balanced accuracy plus 0.20 times group accuracy, minus 0.05 times mean cached deployment seconds exceeding one second. A disclosed complexity proxy additionally subtracts 0.002 log₂(1 + native head count) and 0.002 per extra structural stage. This proxy is not an objective novelty measure. The revised cohort and cleanup require fresh quality scoring; prior frozen predictions are reused only because training, recipes and image inputs are byte-identical. Partial earlier scores are not pooled. No result from an unfinished generation is treated as the final best method.')
    heading(doc,'4 4 Metrics timing and statistical interpretation',2)
    para(doc,'Pixel macro IoU averages positive tissue roles with nonzero union on valid annotated pixels, including spurious predicted roles. Unannotated or unmapped reference pixels are ignored, while annotated tissue lost through foreground estimation remains error. Each image contributes equally to image-mean IoU. Group accuracy compares one prediction with the hidden source identity; balanced accuracy averages recalls of observed reference roles. These endpoints measure different errors and are reported separately.')
    para(doc,'Chiamaka uses an Intel Core i7-13700 CPU, two numerical threads and low-priority workers. Controlled refinement times are medians of three repeats after warm-up and include guide construction, solves and native projection. They exclude encoder inference, RGB construction and disk loading. Warm reference classification timing excludes startup, pooling, context construction and polygon export. In the evolving study, the recorded cached deployment sum includes AF geometry, pooling, required descriptors, warm prediction, export, compact fitting and final cleanup; repeated identical vectors retain their measured charge even when retrieved from cache. Neither timing protocol establishes cold end-to-end or complete-WSI throughput.')
    para(doc,'Images have been inspected and test results adaptively reused throughout development. The 51-image reference and the 116-image expansion are development evaluations, not untouched tests. Changes in cohort composition prevent interpreting a change in the aggregate score as an algorithmic gain. The source-group bootstrap does not remedy repeated test inspection, patient-link uncertainty or annotation-informed field selection.')
    heading(doc,'5 Results')
    heading(doc,'5 1 Conditional refinement',2)
    table(doc,'Table 2. Completed matched conditional-refinement results on 39 development ROIs. Runtime is the measured refinement stage; encoder and RGB costs are excluded.',
        ['Method','Mean macro IoU','Seconds per ROI'],[[r['name'],f"{r['mean_iou']:.4f}",f"{r['mean_seconds']:.4f}"] for r in controlled['ranking']],[3.3,1.65,1.7])
    para(doc,'Bora-AF has image-mean IoU 0.6628 versus 0.5862 for the input tiles and 0.6435 for the convergence-qualified bilateral reference. The paired source-group mean difference from the bilateral reference is 0.0208, with a development bootstrap interval of 0.0093 to 0.0337. Its mean measured refinement time is 0.1676 seconds; the bilateral notebook requires 0.4569 seconds. The soft graph control is faster at 0.0434 seconds but has lower mean IoU, 0.6061. These observations establish an accuracy–time tradeoff within this experiment, not universal speed superiority.')
    para(doc,'Bora-VAF yields 0.6621 IoU in 0.3643 seconds. Its paired group-mean difference from AF is −0.00039, with interval −0.00215 to 0.00151. The optional variance model therefore has no demonstrated gain in this cohort and is not selected as the current core.')
    figure(doc,ROOT/'figures/controlled.png','Figure 1. Matched conditional refinement. A shows native-mask macro IoU; B shows the separately measured refinement cost. All methods share inputs and quality masks. Bilateral timings refer to the tested convergence-qualified notebook implementation, not its authors’ optimized implementation. These 39 ROIs were used during development.')
    heading(doc,'5 2 Neighbor context on the PUMA free reference cohort',2)
    labels={'R01':'Tile evidence vote','R02':'Pooled mean','R03':'Mean and dispersion','R04':'Composition','R05':'Neighbor context'}
    table(doc,'Table 3. Recalculated reference results on the same 51 development test images and 173 whole groups. Head time is warm native prediction only.',
        ['Method','Group accuracy','Balanced accuracy','Final IoU','Head ms'],
        [[labels[m['family']],f"{100*m['top1']:.1f}%",f"{100*m['balanced_accuracy']:.1f}%",f"{m['mean_geojson_iou']:.4f}",f"{1000*m['mean_head_seconds']:.3f}"] for m in view['models']],
        [2.12,1.15,1.25,1.1,1.03])
    para(doc,'Neighbor context correctly identifies 158 of 173 groups, with group accuracy 0.9133, balanced accuracy 0.7038 and final GeoJSON IoU 0.6065. Pooled mean reaches 0.8844, 0.6532 and 0.5828, respectively. The image-mean IoU difference is 0.0237. After equal source-group averaging, the difference is 0.0179 and its 95% bootstrap interval spans −0.0031 to 0.0411. Thus the current data suggest a benefit, but do not establish a reliable generalization advantage from context. The 0.163-millisecond mean warm head time does not represent the complete workflow; the saved reference startup refit alone takes approximately 1.69 seconds.')
    para(doc,'Final IoU differs by source: 0.5020 for BCSS, 0.7500 for TIGER, 0.7772 for GCSS and 0.4967 for IGNITE. The latter two estimates use only five and three images and cannot characterize their complete releases. The expanded evaluation increases IGNITE and GCSS test coverage, while preserving all previous internal split assignments. Its final comparative scores remain pending.')
    figure(doc,ROOT/'figures/examples.png', '. '.join(['Figure 2. Whole-group refinement and identification examples',f'Top two rows {examples[0]} show the median Neighbor context IoU example among the 51 retained cases; bottom two rows {examples[1]} show the lowest-IoU example','Expert, anonymous tile and final tissue overlays are shown over H&E; RGB is the unchanged evidence conversion','Every panel has a 100 µm scale bar','Colors in expert and prediction panels follow the displayed tissue legend; anonymous tile colors denote group IDs','Examples are chosen descriptively from inspected results and are not evidence of independent performance'])+'.')
    heading(doc,'5 3 Polygon compactness and final area checks',2)
    para(doc,'Across the retained five methods and 51 images, 255 exports contain 1,801,143 boundary segments before compact fitting and 725,795 afterward, a 59.7% reduction. The latest continuous-area cleanup removes 709 small polygon parts and fills 50 small polygon holes that were not fully controlled by raster-only checking. Independent audits cover 989,643,305 native output pixels and find no positive polygon part, polygon hole, positive raster region or enclosed raster hole below 20, with zero measured interclass overlap. These figures describe all five outputs, not 255 independent images or patients. Compactness is an engineering endpoint; its effect on anatomy still requires pathology review and task-specific morphometric validation.')
    heading(doc,'6 Discussion')
    para(doc,'The completed evidence supports Bora-AF as a practical conditional refiner of coarse tissue groups. Regularizing class appearances rather than the final label plane allows groups to compete throughout the supported image, avoiding a fixed edit band. The low-dimensional field grid makes the solves inexpensive relative to dense feature extraction. The result remains contingent on the supplied coarse information, foreground mask and ability of the evidence representation to distinguish compartments. A group missing completely from the coarse input cannot be recovered as a new anonymous identity.')
    para(doc,'Neighbor context provides a transparent way to use the tissue relationship around a group without training a message-passing graph network. Shared physical boundary length controls which neighbors influence the contrast, and one native head handles the augmented vector. The inspected reference results favor this choice, but the paired interval crosses zero and the benchmark has been reused adaptively. Context may also encode dataset-specific compartment relationships. A validation experiment holding out institutions, organs and donor groups is required to establish transportability. The classifier family should remain provisional until the ongoing development search has been completed and a separate final evaluation has been frozen.')
    para(doc,'The anonymous task must be interpreted carefully. Removing class names from an expert-derived coarse map does not remove its expert information: group count, assignment and geometry still encode the tissue partition. The RGB projection is also conditioned on those coarse groups. Results therefore assess conditional geometric recovery plus tissue identification, not annotation-free discovery or performance of an autonomous dense segmenter. A realistic upstream prediction experiment, with tile-class mistakes and omitted classes, is a necessary extension.')
    para(doc,'The vector stage addresses a practical issue that raster scores alone miss. Smoothing can leave subpixel polygon slivers whose area is below the intended threshold even though they touch enough pixel centers to pass a native component test. Shared arcs reduce staircase edges consistently between neighbors, while continuous-area and raster checks enforce complementary constraints. By definition, the 20-native-pixel cutoff has different physical area at different pixel spacings. It should be distinguished from the physically calibrated chord tolerance and tested for sensitivity before use in small-structure morphometry.')
    para(doc,'The principal methodological uncertainty is incremental novelty relative to graph smoothers, local appearance models and contextual regional classifiers. A stronger claim would require an explicit ablation isolating appearance-field fitting from H&E conductance, global statistics and evidence conversion, together with contemporary dense students and feature upsampling methods under equal input and computational budgets. More evolutionary candidates cannot replace independent validation. Model selection on adaptively reused test images yields development evidence, even if each generation freezes its predictions before scoring.')
    para(doc,'Other limitations include boundary-enriched acquisition, possible unlinked donors across sources, overlapping ontology levels, rare or untrained roles, uncertain optical resolving power beyond the recorded sampling, and the coupling of disconnected pieces within a group. False background removal can erase pale tissue; such errors remain in annotated quality scores. Native-coordinate polygons and ROI latency do not establish complete-slide accuracy, memory scaling or throughput. Clinical utility, observer agreement, uncertainty calibration and downstream measurement reliability have not yet been demonstrated.')
    heading(doc,'7 Conclusion')
    para(doc,'Bora-AF combines physically calibrated appearance fields, conditional multiclass competition and compact shared-border export. Neighbor context supplies a simple native fastPLS classifier for whole refined anonymous groups. Completed development experiments show useful refinement and identification performance, while the revised multiclass cohort broadens the ongoing evaluation. Claims about autonomous segmentation, independent clinical generalization or best-in-class deployment efficiency require additional evidence.')
    heading(doc,'Data code and declarations')
    para(doc,'Bora is maintained at https://github.com/tkcaccia/Bora. The published coverage converter is identified by commit 08b63a6595df687bd20a516f43f65dc531e664dd; later appearance, group-classification and compact-cleanup code is identified by run-time file hashes. A single archived public release of the complete revised pipeline remains to be prepared. Dataset provider records and acquisition inventories retain image/mask checksums, crop coordinates, licensing, annotation mappings, WSI/ROI type and per-axis resolution. The derived arrays and local viewers are not represented as an already public archival data release. fastPLS 0.3.0 and its exact core and binding revisions are recorded in the model manifests [7].')
    para(doc,'This research uses previously released, deidentified public data under provider-specific terms. The authors must confirm the institutional determination for public-data reuse, funding, competing interests, contributions, correspondence and final affiliations before submission. No clinical deployment, new patient recruitment, funding source or ethics approval identifier is asserted here. The author list is retained from the supplied manuscript. This is an unsubmitted working draft; the ongoing ten-generation study is not reported as completed.')
    heading(doc,'References')
    refs=[
        ('Zimmermann E, Vorontsov E, Viret J, et al. Virchow2 Scaling Self Supervised Mixed Magnification Models in Pathology. 2024. arXiv 2408.00738.','https://arxiv.org/abs/2408.00738'),
        ('Grady L. Random walks for image segmentation. IEEE Transactions on Pattern Analysis and Machine Intelligence. 2006;28:1768–1783.','https://doi.org/10.1109/TPAMI.2006.233'),
        ('Barron JT, Poole B. The Fast Bilateral Solver. ECCV. 2016.','https://arxiv.org/abs/1511.03296'),
        ('Wang L, He L, Mishra A, Li C. Active contours driven by local Gaussian distribution fitting energy. Signal Processing. 2009;89:2435–2447.','https://doi.org/10.1016/j.sigpro.2009.03.014'),
        ('Fu S, Hamilton M, Brandt L, Feldman A, Zhang Z, Freeman WT. FeatUp A Model Agnostic Framework for Features at Any Resolution. ICLR. 2024.','https://arxiv.org/abs/2403.10516'),
        ('Lv J, Zhu Y, Okada S, Renna MS, Goldin R, Antonowicz S, Raza SEA. Computationally Efficient Pathology Segmentation using Knowledge Distillation from Foundation Models. 2026. arXiv 2609.03947v1. Preprint supplied by the author.','https://arxiv.org/abs/2609.03947'),
        ('Cacciatore S. fastPLS py native C++ numerical core and Python bindings. Version 0.3.0; binding fc0c011969b754ee196aacc220a217783afeb05d; core 82bbc48a0d69e4bd0d7c261fdcc8e636694133b1.','https://github.com/tkcaccia/fastPLS-py'),
        ('Amgad M, Elfandy H, Hussein H, et al. Structured crowdsourcing enables convolutional segmentation of histology images. Bioinformatics. 2019;35:3461–3467.','https://doi.org/10.1093/bioinformatics/btz083'),
        ('van Rijthoven M, Aswolinskiy W, Tessier L, et al. TIGER training dataset ROI-level annotations of WSIROIS subset. Zenodo. 2022.','https://zenodo.org/records/6014422'),
        ('IGNITE data toolkit. Tissue and cell-level annotated H&E and PD-L1 histopathology image dataset in non-small cell lung cancer. Zenodo record 17735903.','https://zenodo.org/records/17735903'),
        ('GCSS gastric tissue segmentation data accompanying Semi-Supervised Pixel Contrastive Learning Framework for Tissue Segmentation in Histopathological Image. Zenodo record 8176354.','https://zenodo.org/records/8176354'),
        ('CAMELYON16. Original whole-slide images and tumor annotations. Registry of Open Data on AWS.','https://registry.opendata.aws/camelyon/'),
        ('Lems CM, et al. BEETLE multicentric breast tissue segmentation dataset. Zenodo record 16812932.','https://zenodo.org/records/16812932'),
        ('OCELOT. Challenge data and author repository.','https://github.com/lunit-io/ocelot23algo'),
        ('PESO test release. Version 1.1. Zenodo record 5137717. Reused as development data in this study.','https://zenodo.org/records/5137717'),
        ('PUMA. Released training tissue annotations and source image context. Zenodo record 15050523.','https://zenodo.org/records/15050523'),
        ('GleasonXAI. TissueArray images and region annotations. Figshare version 1.','https://doi.org/10.6084/m9.figshare.27301845.v1')]
    for i,(text,url) in enumerate(refs,1):reference(doc,i,text,url)
    heading(doc,'Appendix A Prospective applicability study').paragraph_format.page_break_before=True
    para(doc,'This protocol is proposed and has not yet generated results. The completed benchmark in Tables 1–3 evaluates conditional refinement and tissue identification. The next experiment will test whether the same frozen Bora operator improves independently generated coarse maps. The three input conditions will be reported separately: expert-derived tiles, Virchow2 feature clusters, and source-trained fastPLS tile predictions. They answer different questions and their raw scores will not be pooled as if the inputs carried equal information.')
    para(doc,'For the supervised example, one annotated source WSI will supply training tile features and expert tissue targets. The source will be selected for annotation coverage and required class representation without inspecting target performance. Native fastPLS will use physically calibrated Virchow2 tile summaries, balanced training sampling and source-only spatial validation. Its component count, class vocabulary, confidence handling and tile spacing will freeze before inference on a different patient’s target WSI. Target annotations will remain inaccessible to preprocessing, feature selection, coarse prediction and Bora. They will be decoded only after both the raw tile output and refined output have been frozen. This is transfer from one annotated slide, not one-shot learning: one WSI may provide thousands of labeled training tiles.')
    para(doc,'For the clustering condition, anonymous groups will be produced from target Virchow2 features without expert tile assignments or a target-informed number of groups. Feature normalization, clustering rules, group-count selection and RGB construction will be fixed using development sources. Bora will refine that anonymous partition. Its primary assessment will use label-invariant grouping and boundary measures. If named tissue outputs are required, a Neighbor context readout must be trained only on permitted source material and frozen before target scoring; target annotation matching must not be used to assign semantic names. The source-trained tile predictor and the group identifier are different model roles.')
    para(doc,'The initial qualitative demonstration can use existing full BEETLE WSIs with pixel annotations on Chiamaka, at verified image and annotation pyramid levels. Those slides have already contributed inspected development windows, so this first example cannot be described as independent external validation. BEETLE labels such as other tissue will retain their provider meaning and will not be renamed stroma. A publication evaluation should repeat transfers across several source slides and unseen target patients, preferably including a held-out institution, with all learned components retrained without target donors. Classes absent from the source will be reported explicitly rather than removed after seeing target errors.')
    para(doc,'For each source–target pairing, the main comparison is the exact same coarse map before and after Bora. The fixed RGB9+19 formulas may use the independently generated coarse groups, but no expert-derived RGB cache from the controlled benchmark may enter these arms. All tissue will be encoded independently of annotation coverage. Outputs will show original H&E, expert annotation, raw tiles or clusters, evidence RGB and the final simplified GeoJSON mask. Pixel macro IoU and Dice, physical boundary F1 and surface distance, topology and geometric area error will accompany full-pipeline wall time, peak memory and time per tissue area. Confidence intervals will resample patients or source slides, not individual tiles.')
    para(doc,'The MedIA-oriented validation will pair these applicability arms with matched image-guided refiners and a compact dense or distilled segmentation baseline, and with minimal ablations of the AF appearance model, H&E conductance, context and RGB evidence. The present 116-image search remains development work. A separate evaluation cohort and the complete operating point must freeze before final testing. A single successful transfer is an illustration; consistent paired improvements across independent cases, a clearly isolated methodological contribution and honest deployment cost are the evidence needed for the stronger journal claim.')
    doc.core_properties.title='Bora appearance fields and Neighbor context'
    doc.core_properties.author='Bora research team';doc.core_properties.subject='Working article draft with controlled refinement and whole anonymous group identification'
    path=ROOT/'Bora_Neighbor_Context_manuscript.docx';doc.save(path);return path


def review():
    doc=new_document('Critical review and journal strategy for Bora',
        'Simulated Medical Image Analysis reviewer assessment of the revised working draft  |  This is an internal assessment, not an actual editorial decision')
    doc.styles['Normal'].paragraph_format.line_spacing=1.05
    doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    heading(doc,'Editorial recommendation')
    para(doc,'I would not recommend submission to Medical Image Analysis in its present evidential state. The manuscript describes a coherent and useful conditional refinement workflow, but its strongest quantitative evidence comes from inspected, reference-conditioned development data and limited classical comparators. The mathematical and contextual components are individually established. My simulated recommendation is rejection with encouragement to resubmit after substantially stronger validation, rather than a cosmetic revision. This assessment does not predict an actual editor’s decision.')
    heading(doc,'What is convincing')
    para(doc,'The revised paper now defines the task correctly: retain the previous expert-derived tile assignments, anonymize names, refine geometry with AF, and identify whole groups using native PLS. The AF field system is explicit and its fixed-observation convexity claim is appropriately limited. The paper separates warm head latency from encoder and geometry cost, distinguishes ROIs from complete slides, and measures the rasterization of the actual final vectors. The paired interval for the bilateral comparison is useful. Showing the worst reference example and retaining rare-class errors improves credibility.')
    heading(doc,'Major concerns that determine journal readiness')
    heading(doc,'1 The current task does not establish autonomous segmentation',2)
    para(doc,'The coarse groups and RGB conversion inherit expert annotation information. Hiding numeric class names does not remove class count or partition information. The 91.3% identification accuracy should not be presented as an independent segmentation result. The revision states this clearly, but the scientific validation still needs a realistic upstream coarse predictor or separately generated unsupervised groups on unseen cases. Include an explicit annotation-dependence diagram and test missing groups, wrong tile labels and incorrect group counts. Analyze this experiment separately from controlled quantization recovery.')
    heading(doc,'2 Repeated test inspection prevents a final generalization claim',2)
    para(doc,'The ten-generation search deliberately uses validation and formerly reserved test scores. It is a development search. Freezing each generation before scoring prevents accidental after-the-fact file changes but does not protect against adaptive test overfitting. Before the final method is selected, reserve a new institution or patient cohort that has never informed the search. Freeze the complete algorithm, preprocessing, ontology, thresholds and operating point before evaluating it. Do not recover an untouched test by simply renaming the existing validation or test split. Verify patient links across BCSS and TIGER.')
    heading(doc,'3 The mathematical novelty remains vulnerable',2)
    para(doc,'AF solves a regularized quadratic appearance model on an H&E-guided graph; local Gaussian models, confidence-weighted graph smoothers and image-guided regularization already exist. Neighbor context concatenates a group mean with a boundary-weighted neighborhood contrast and then applies established PLS-LDA. Compact shared-border simplification is valuable engineering but does not alone establish a new segmentation principle. The manuscript needs a precise closest-method comparison and an ablation showing which modeling decision cannot be replaced by a simpler equivalent at the same cost. Avoid claiming that combining established operations creates novelty by itself.')
    heading(doc,'4 Contemporary comparators and matched budgets are missing',2)
    para(doc,'The bilateral reference notebook is not its optimized implementation, and the current graph control is not the original random walker. Add a compact supervised dense segmentation baseline, preferably a published distilled student or a carefully trained efficient U-Net/nnU-Net variant, on identical training sources and output labels. Also compare a strong feature upsampling route and a direct label-regularization control under matched coarse input, H&E access and hardware budgets. Separate the conditional-refinement question from autonomous dense segmentation. Do not compare a cached CPU head with cold encoder-plus-decoder inference.')
    heading(doc,'5 Neighbor context is promising but not yet established',2)
    para(doc,'Across 51 non-PUMA images, context improves image-mean final IoU from 0.5828 to 0.6065 and balanced accuracy from 65.3% to 70.4%. However, the source-group bootstrap interval for the IoU difference is −0.0031 to 0.0411. This crosses zero before accounting for adaptive selection. Keep the family as the user-chosen provisional description, not a confirmed winner. Compare it with the same native PLS head on mean alone; neighborhood averaging alone; shuffled adjacency; physical versus unweighted edges; and a modest graph model. Perform these ablations on development data, then freeze the result.')
    heading(doc,'6 Ontology and biological interpretation require expert review',2)
    para(doc,'Twenty-eight mapped roles combine compartments, cell populations, organ-specific structures and grades. Only 25 have represented hard-target training classes. Report role frequencies, confusion matrices, untrained-role errors and dataset-specific mappings. A group can contain disconnected pieces or mixed tissue after refinement. A pathologist should review both successful and failed groups, and assess whether a shared identity is defensible. Establish whether context exploits organ or provider identity rather than transferable tissue relationships.')
    heading(doc,'7 Geometry fidelity needs a biological endpoint',2)
    para(doc,'Zero sub-20 fragments and zero overlap establish implementation compliance, not preservation of clinically relevant anatomy. The 20-native-pixel rule corresponds to different physical areas at different spacings; the earlier 500-pixel converter filter must remain clearly distinguished. Test tolerance and area thresholds against boundary F1 in micrometres, surface distance, topology, small-gland or vessel retention and downstream measurements. Include blinded expert review of diagonal smoothing. Give both before/after annotation scores; agreement with the preceding predicted geometry is not agreement with an expert.')
    heading(doc,'8 Efficiency and reproducibility need end to end evidence',2)
    para(doc,'Report cold and warm wall-clock costs with and without cached Virchow2 features, peak memory, encoder hardware, WSI size, total processed tissue area and physical throughput. Include the exact costs of polygon fitting and cleanup. Retain the above-one-second penalty as an operating preference, but report the unpenalized accuracy–time Pareto front as well. Archive a single reproducible release, pinned native library and runnable protocols rather than relying on dirty-worktree hashes and local viewers. Dataset licenses and author declarations must be finalized.')
    heading(doc,'Revisions already incorporated into the article')
    para(doc,'The revised text removes obsolete random KMeans and connected-fragment claims; describes the full Neighbor context feature equation and native head; labels reference-conditioned inputs and adaptive reuse; excludes PUMA from current evaluation; separates the completed 51-image results from the ongoing 116-image study; reports the confidence interval crossing zero; documents the native versus physical geometry thresholds; distinguishes the 500-pixel converter filter from final20 cleanup; and avoids equating native coordinates or millisecond heads with full-slide speed. The original author list and original source document remain preserved.')
    heading(doc,'Journal strategy')
    table(doc,'Journal fit is an editorial judgment based on the current scope and evidence, not a promised acceptance or a ranking by impact factor.',
        ['Journal','Fit for this work','Current recommendation'],[
            ['Medical Image Analysis','Strong methodological scope match for image analysis and segmentation','Ambitious principal target after independent validation and stronger methodological evidence'],
            ['IEEE Transactions on Medical Imaging','Strong scope match for novel mathematical imaging methods','Equally demanding alternative; not an easier fallback for the same validation gaps'],
            ['Computerized Medical Imaging and Graphics','Good match for a rigorously validated computational pathology workflow and visualization pipeline','Most realistic specialist target if the final contribution remains primarily a useful efficient workflow']],[2.05,2.3,2.3])
    para(doc,'My preferred highest-ambition target is Medical Image Analysis if the paper demonstrates a substantive appearance-field contribution and generalization beyond expert-conditioned development inputs. The publisher explicitly includes segmentation, geometric/statistical models and tissue-scale microscopy in its scope [1]. IEEE TMI similarly welcomes new analysis methods, but its stated criteria emphasize methodological novelty and impact beyond incremental advancement [2,3]. At the current contribution level, Computerized Medical Imaging and Graphics is the more realistic fit; its official scope explicitly includes computational methods, visualization, digital pathology and optical microscopy [4]. These fit judgments are my inference from the manuscript and official scopes, not statements from the journals.')
    para(doc,'I would not position the current study for a general clinical or translational flagship simply on the basis of ROI accuracy and speed. Such positioning would need a substantial new biological or clinical result, external evidence and a compelling translational use case. Increasing the number of classifiers alone would not supply that evidence. No acceptance probabilities, impact-factor claims or submission-fee assumptions are made here.')
    heading(doc,'Prioritized route to a stronger submission')
    for text in [
        'Finish the already authorized development search and freeze one method. Retain Neighbor context in the present draft until the full comparison is complete.',
        'Audit the ontology and all source-group links; create a genuinely untouched external cohort and preserve it from all evolutionary selection.',
        'Add a realistic upstream coarse-map experiment and matched contemporary dense or distilled baselines, with annotation dependence explicitly controlled.',
        'Complete the minimal AF/context ablations and report source-group confidence intervals, per-role failures, physical boundary metrics and topology.',
        'Measure full-pipeline and complete-slide cost, then archive the final code, models, protocols and data manifests.',
        'Choose the journal from the resulting scientific contribution: MedIA for a clear new method with convincing independent evidence, CMIG for a strong validated efficient pathology workflow.']:
        para(doc,text,'List Number')
    heading(doc,'Sources for journal scope')
    refs=[('Elsevier. Medical Image Analysis official journal description. Accessed18 October2026.','https://shop.elsevier.com/journals/medical-image-analysis/1361-8415'),
        ('IEEE TMI. Scope. Accessed18 October2026.','https://ieeetmi.org/scope/'),
        ('IEEE TMI. Key Criteria for Publication. Accessed18 October2026.','https://ieeetmi.org/key-criteria-for-publication/'),
        ('Elsevier. Computerized Medical Imaging and Graphics official journal description. Accessed18 October2026.','https://shop.elsevier.com/journals/computerized-medical-imaging-and-graphics/0895-6111')]
    for i,(text,url) in enumerate(refs,1):reference(doc,i,text,url)
    doc.core_properties.title='Bora critical review and journal strategy';doc.core_properties.author='Internal scientific assessment'
    path=ROOT/'Bora_MedIA_critical_review_and_journal_strategy.docx';doc.save(path);return path


if __name__=='__main__':
    source_hash=sha(SOURCE);examples=figures();manuscript=article(examples);assessment=review()
    inputs=[SOURCE,ROOT/'evidence/REGISTRY.json',ROOT/'evidence/REFERENCE_SUMMARY.json',GROUP/'viewer-no-puma-area20-v7/data.json',GROUP/'geojson-no-puma-area20-v7/VALIDATION.json',EVO/'COHORT_REVISION.json',REPORTS/'bora-supervised-2026-10-18-v1/CONTROLLED_FINAL.json']
    provenance=dict(provisional_classifier='R05-c48 Neighbor context',prediction_unit='Whole anonymous group including disconnected pieces',
        reference_images=51,reference_groups=173,expanded_test_images=78,expanded_validation_images=38,
        expanded_results_included=False,ten_cycle_results_complete=False,source_unchanged=sha(SOURCE)==source_hash,
        source_sha256={str(p):sha(p) for p in inputs},example_selection=dict(median=examples[0],lowest_iou=examples[1]),
        outputs={str(p):sha(p) for p in [manuscript,assessment]},unsubmitted_draft=True)
    (ROOT/'MANUSCRIPT_PROVENANCE.json').write_text(json.dumps(provenance,indent=2)+'\n')
    print(manuscript);print(assessment)
